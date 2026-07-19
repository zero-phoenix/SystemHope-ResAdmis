#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Script de validación: Verifica que un archivo .docx sea auténtico (ZIP OOXML válido).

Uso:
    python validar_docx.py archivo.docx
    python validar_docx.py "ruta/con espacios/archivo.docx"
"""

import sys
import zipfile
from pathlib import Path
from docx import Document


def validar_docx(ruta):
    """
    Valida que un archivo .docx sea auténtico (ZIP OOXML).

    Retorna:
        (bool, str): (es_válido, mensaje)
    """
    path = Path(ruta)

    # Verificación 0: Existe
    if not path.exists():
        return False, f"Archivo no encontrado: {path}"

    # Verificación 1: Es un archivo ZIP válido
    try:
        with zipfile.ZipFile(path, 'r') as z:
            files = z.namelist()
    except zipfile.BadZipFile:
        return False, "ERROR: No es un archivo ZIP válido. Probablemente es un archivo de texto plano con extensión .docx (CORRUPTO)"
    except Exception as e:
        return False, f"ERROR al leer ZIP: {e}"

    # Verificación 2: Contiene archivos OOXML requeridos
    required_files = ['[Content_Types].xml', 'word/document.xml', '_rels/.rels']
    missing = [f for f in required_files if f not in files]

    if missing:
        return False, f"CORRUPTO: Faltan archivos OOXML requeridos: {', '.join(missing)}"

    # Verificación 3: Puede abrirse con python-docx
    try:
        doc = Document(str(path))
        para_count = len(doc.paragraphs)
        table_count = len(doc.tables)
    except Exception as e:
        return False, f"ERROR al abrir con python-docx: {e}"

    # Verificación 4: Tamaño mínimo (archivos válidos > 1 KB)
    size_kb = path.stat().st_size / 1024
    if size_kb < 1:
        return False, f"ALERTA: Archivo demasiado pequeño ({size_kb:.1f} KB). Podría estar vacío."

    # Todas las verificaciones pasaron
    info = f"VALIDO [{size_kb:.1f} KB] - {para_count} parrafos, {table_count} tablas, {len(files)} archivos internos"
    return True, info


def main():
    if len(sys.argv) < 2:
        print("Uso: python validar_docx.py <ruta_al_archivo.docx>")
        print("\nEjemplos:")
        print('  python validar_docx.py mi_resolucion.docx')
        print('  python validar_docx.py "D:\\\\BETTER CALL DAVID\\\\ResAdmi\\\\productos (resoluciones) elaborada por claude\\\\ADM.docx"')
        sys.exit(1)

    ruta = sys.argv[1]
    valido, mensaje = validar_docx(ruta)

    # Imprimir resultado
    symbol = "[OK]" if valido else "[ERROR]"
    print(f"\n{symbol} {Path(ruta).name}")
    print(f"    {mensaje}\n")

    # Exit code
    sys.exit(0 if valido else 1)


if __name__ == '__main__':
    main()
