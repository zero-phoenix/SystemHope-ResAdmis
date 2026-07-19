import docx
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
import copy
import sys
import os

def clean_run(run):
    run.font.name = 'Arial Narrow'
    run.font.size = Pt(11)
    run.font.highlight_color = None
    run.font.color.rgb = RGBColor(0, 0, 0)
    rPr = run._element.get_or_add_rPr()
    shd = rPr.find(qn('w:shd'))
    if shd is not None:
        rPr.remove(shd)
    highlight = rPr.find(qn('w:highlight'))
    if highlight is not None:
        rPr.remove(highlight)

def clean_paragraph_shading(p):
    pPr = p._element.get_or_add_pPr()
    shd = pPr.find(qn('w:shd'))
    if shd is not None:
        pPr.remove(shd)
    # also remove paragraph level highlight if any (rare but possible)
    highlight = pPr.find(qn('w:highlight'))
    if highlight is not None:
        pPr.remove(highlight)

def remove_numbering(p):
    pPr = p._element.get_or_add_pPr()
    numPr = pPr.find(qn('w:numPr'))
    if numPr is not None:
        pPr.remove(numPr)

def copy_pPr(source_p, target_p):
    source_pPr = source_p._element.pPr
    if source_pPr is not None:
        target_pPr = target_p._element.get_or_add_pPr()
        target_p._element.replace(target_pPr, copy.deepcopy(source_pPr))

def add_run(p, text, bold=False):
    run = p.add_run(text)
    run.bold = bold
    clean_run(run)
    return run

doc = docx.Document(r'D:\CC1\Modelos al 30-06-26\MODELO xxxx exp. ADM ver. David 1 DDO,  VARIAS IMPUTACIONES.docx')

expediente = "2142-2026/CC1"
denunciante = "EFRAÍN CORONEL QUISPE (SEÑOR CORONEL)"
denunciado = "LA POSITIVA SEGUROS Y REASEGUROS S.A.A. (LA POSITIVA)"
fecha_res = "Lima, 17 de julio de 2026"
fecha_denuncia = "20 de abril de 2026"

hechos = [
    "(i)\tEl 06 de julio de 2023, contrató un seguro vehicular con la empresa La Positiva Seguros y Reaseguros S.A.A., respecto del vehículo de placa de rodaje BVO-787.",
    "(ii)\tEl 22 de mayo de 2024, el vehículo fue objeto de robo total mientras se encontraba estacionado en el frontis del inmueble donde era guardado.",
    "(iii)\tMediante carta LPG.SINIESTROS.RCH IR 000000853 - 2024 de fecha 25 de junio de 2024, la aseguradora comunicó el rechazo de la cobertura del siniestro, sustentando su decisión en una supuesta utilización del vehículo en forma de alquiler.",
    "(iv)\tCon posterioridad, la aseguradora continuó realizando cobros a la tarjeta CMR por un periodo adicional hasta julio de 2025 por concepto de renovación del seguro que vencía en julio de 2024, no obstante tener conocimiento del robo. Asimismo, le envió correos exigiendo pagos por otro vehículo asegurado.",
    "(v)\tLa aseguradora no brindó información clara ni entregó la documentación adecuada relacionada con la póliza contratada."
]

medida_correctiva = "2.\tEl señor Coronel solicitó, en calidad de medida correctiva, que la aseguradora cumpla con otorgar la cobertura del Seguro Vehicular por el siniestro de robo total ocurrido el 22 de mayo de 2024, el pago de la suma asegurada conforme a la póliza ascendente a USD 12 450,00, y la devolución de los cobros indebidos. Asimismo, requirió el reembolso de costas y costos del presente procedimiento."

imputaciones_analisis = [
    "La Secretaría Técnica de la Comisión de Protección al Consumidor 1 (en adelante, la Secretaría Técnica), en ejercicio de sus facultades, considera que el hecho denunciado, consistente en que la compañía aseguradora habría negado injustificadamente al denunciante la cobertura del Seguro Vehicular (Póliza N° 230245150) para el vehículo con placa de rodaje BVO-787 por el siniestro de robo ocurrido el 22 de mayo de 2024, rechazo comunicado mediante carta LPG.SINIESTROS.RCH IR 000000853 - 2024 de fecha 25 de junio de 2024; involucraría una presunta afectación a sus expectativas, quien no habría encontrado una correspondencia entre lo que esperaba recibir de parte del proveedor y lo que realmente recibió. Por consiguiente, corresponde calificar el hecho materia de denuncia como una presunta infracción al deber de idoneidad, tipificado en los artículos 18° y 19° del Código.__F1__",
    "Así también, la Secretaría Técnica, en ejercicio de sus facultades, considera que el hecho denunciado, consistente en que la compañía aseguradora no habría cumplido con entregar de manera adecuada las condiciones generales y particulares de la póliza del Seguro Vehicular (Póliza N° 230245150). Por consiguiente, corresponde calificar el hecho materia de denuncia como una presunta infracción al deber de entregar a los usuarios copia de los contratos y demás documentación relacionada con dichos actos jurídicos, tipificado en el literal e) del artículo 47° del Código.",
    "Asimismo, la Secretaría Técnica, en ejercicio de sus facultades, considera que el hecho denunciado, consistente en que la compañía aseguradora habría efectuado cobros no aceptados en la tarjeta CMR por concepto de una renovación automática no solicitada del Seguro Vehicular (Póliza N° 230245150), y exigido pagos por otro vehículo asegurado; involucraría una presunta afectación a sus expectativas, quien no habría encontrado una correspondencia entre lo que esperaba recibir de parte del proveedor y lo que realmente recibió. Por consiguiente, corresponde calificar el hecho materia de denuncia como una presunta infracción al deber de idoneidad, tipificado en los artículos 18° y 19° del Código."
]

req_info = "A efectos de tener mayores elementos que sirvan para la resolución definitiva del presente caso, la Secretaría Técnica, en ejercicio de las facultades que la ley le confiere, conviene requerir a la compañía aseguradora que, en un plazo no mayor de cinco (5) días hábiles contado a partir del día siguiente de la notificación de la presente resolución, cumpla con lo siguiente: (i) presentar una copia completa, legible y debidamente suscrita de la Póliza N° 230245150 del Seguro Vehicular; y, (ii) presentar los medios probatorios que acrediten el estado de cuenta y los sustentos de los cobros realizados en la tarjeta CMR del denunciante, así como del otro vehículo asegurado, durante el periodo de julio de 2024 a julio de 2025."

imputaciones_res = [
    "(i)\tPresunta infracción a los artículos 18° y 19° de la Ley N° 29571, Código de Protección y Defensa del Consumidor, en tanto la compañía aseguradora habría negado injustificadamente al denunciante la cobertura del Seguro Vehicular (Póliza N° 230245150) para el vehículo con placa de rodaje BVO-787 por el siniestro de robo ocurrido el 22 de mayo de 2024, rechazo comunicado mediante carta LPG.SINIESTROS.RCH IR 000000853 - 2024 de fecha 25 de junio de 2024.",
    "(ii)\tPresunta infracción al literal e) del artículo 47° de la Ley N° 29571, Código de Protección y Defensa del Consumidor, en tanto la compañía aseguradora no habría cumplido con entregar de manera adecuada las condiciones generales y particulares de la póliza del Seguro Vehicular (Póliza N° 230245150).",
    "(iii)\tPresunta infracción a los artículos 18° y 19° de la Ley N° 29571, Código de Protección y Defensa del Consumidor, en tanto la compañía aseguradora habría efectuado cobros no aceptados en la tarjeta CMR por concepto de una renovación automática no solicitada del Seguro Vehicular (Póliza N° 230245150), y exigido pagos por otro vehículo asegurado."
]

for p in doc.paragraphs:
    clean_paragraph_shading(p)
    for run in p.runs:
        clean_run(run)

mode = "header"
hecho_idx = 0
analisis_idx = 0
resuelve_idx = 0

for p in list(doc.paragraphs):
    text = p.text.strip()
    
    if "EXPEDIENTE" in text and "0672-2026/CC1" in text:
        p.text = ""
        add_run(p, f"EXPEDIENTE\t:\t{expediente}", bold=True)
    elif "DENUNCIANTE" in text and "MEDALITH" in text:
        p.text = ""
        add_run(p, f"DENUNCIANTE\t:\t{denunciante}", bold=True)
    elif "DENUNCIADO" in text and "RÍMAC" in text:
        p.text = ""
        add_run(p, f"DENUNCIADO\t:\t{denunciado}", bold=True)
    elif text.startswith("Lima,"):
        p.text = ""
        add_run(p, fecha_res, bold=False)
        p.paragraph_format.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT
    elif text == "HECHOS":
        mode = "hechos_start"
    elif mode == "hechos_start" and text.startswith("Mediante el escrito"):
        p.text = ""
        add_run(p, f"Mediante el escrito de denuncia del {fecha_denuncia}, el señor Coronel denunció a La Positiva por presuntas infracciones a la Ley N° 29571, Código de Protección y Defensa del Consumidor (en adelante, Código), señalando lo siguiente:")
        mode = "hechos_list"
    elif mode == "hechos_list" and (text.startswith("El xx de abril") or text.startswith("El 5 de abril")):
        # We are at the first dummy paragraph in the template AFTER "Mediante el escrito".
        # We must insert the 'hechos' here, before deleting this paragraph, so they appear AFTER "Mediante el escrito".
        if text.startswith("El xx de abril"): # Only insert once!
            for h in hechos:
                np = p.insert_paragraph_before()
                remove_numbering(np)
                np.paragraph_format.left_indent = Inches(0.5)
                np.paragraph_format.first_line_indent = Inches(-0.5)
                np.paragraph_format.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.JUSTIFY
                add_run(np, h)
        p._element.getparent().remove(p._element)
    elif mode == "hechos_list" and text.startswith("La señora Medina"):
        p.text = ""
        remove_numbering(p)
        p.paragraph_format.left_indent = Inches(0)
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.JUSTIFY
        add_run(p, medida_correctiva)
    elif text == "DE LA ADMISIÓN A TRÁMITE DE LA DENUNCIA":
        mode = "analisis"
    elif mode == "analisis" and (text.startswith("La Secretaría Técnica") or text.startswith("Así también") or text.startswith("Asimismo") or text.startswith("Además")):
        if analisis_idx < len(imputaciones_analisis):
            p.text = ""
            add_run(p, imputaciones_analisis[analisis_idx])
            analisis_idx += 1
        else:
            p._element.getparent().remove(p._element)
    elif mode == "analisis" and text.startswith("En tanto la denuncia"):
        mode = "req"
    elif mode == "req" and text.startswith("A efectos de tener mayores elementos"):
        p.text = ""
        add_run(p, req_info)
    elif text == "RESOLUCIÓN DE LA SECRETARÍA TÉCNICA":
        mode = "resuelve"
    elif mode == "resuelve" and text.startswith("PRIMERO"):
        p.text = ""
        add_run(p, "PRIMERO: ", bold=True)
        add_run(p, f"admitir a trámite la denuncia del {fecha_denuncia} interpuesta por el señor Efraín Coronel Quispe contra La Positiva Seguros y Reaseguros S.A.A., por lo siguiente:")
    elif mode == "resuelve" and text.startswith("Presunta infracción"):
        if resuelve_idx < len(imputaciones_res):
            p.text = ""
            remove_numbering(p)
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.5)
            p.paragraph_format.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.JUSTIFY
            add_run(p, imputaciones_res[resuelve_idx])
            resuelve_idx += 1
        else:
            p._element.getparent().remove(p._element)
    elif mode == "resuelve" and text.startswith("SEGUNDO:"):
        p.text = ""
        add_run(p, "SEGUNDO: ", bold=True)
        add_run(p, f"tener por ofrecidos los medios probatorios presentados en el escrito de denuncia del {fecha_denuncia}.")
    elif mode == "resuelve" and text.startswith("TERCERO:"):
        p.text = ""
        add_run(p, "TERCERO: ", bold=True)
        add_run(p, "requerir a La Positiva Seguros y Reaseguros S.A.A. que cumpla con lo siguiente:")
    elif mode == "resuelve" and text.startswith("CUARTO:"):
        p.text = ""
        add_run(p, "CUARTO: ", bold=True)
        add_run(p, f"correr traslado de la denuncia interpuesta el {fecha_denuncia} a La Positiva Seguros y Reaseguros S.A.A., para que, de conformidad con lo dispuesto por el artículo 26° de la Ley sobre Facultades, Normas y Organización del Indecopi, aprobada por Decreto Legislativo 807, presente sus descargos en un plazo no mayor de cinco (5) días hábiles contados desde la notificación.")
    elif mode == "resuelve" and text.startswith("QUINTO:"):
        p.text = ""
        add_run(p, "QUINTO: ", bold=True)
        add_run(p, req_info.replace("A efectos de tener mayores elementos que sirvan para la resolución definitiva del presente caso, la Secretaría Técnica, en ejercicio de las facultades que la ley le confiere, conviene requerir a la compañía aseguradora que, en un plazo no mayor de cinco (5) días hábiles contado a partir del día siguiente de la notificación de la presente resolución, cumpla con lo siguiente:", "requerir que, en un plazo no mayor de cinco (5) días hábiles contado a partir del día siguiente de la notificación de la presente resolución, el proveedor denunciado cumpla con lo siguiente:"))
    elif mode == "resuelve" and text.startswith("DÉCIMO:"):
        p.text = ""
        add_run(p, "DÉCIMO: ", bold=True)
        add_run(p, "requerir al señor Efraín Coronel Quispe para que, dentro del plazo de dos (2) días hábiles siguientes a la fecha en que reciba la notificación en su bandeja de correo electrónico, efectúe la confirmación de recepción de la notificación remitida por este despacho a su correo electrónico, de conformidad con el segundo párrafo del numeral 4 del artículo 20° del Texto Único Ordenado de la Ley del Procedimiento Administrativo General, aprobado mediante Decreto Supremo 006-2026-JUS__F2__, bajo apercibimiento de rehacer el acto de notificación y notificarle conforme al numeral 1 del artículo 20° del citado cuerpo normativo.")
    elif mode == "resuelve" and text.startswith("DÉCIMO PRIMERO:"):
        p.text = ""
        add_run(p, "DÉCIMO PRIMERO: ", bold=True)
        add_run(p, "requerir a La Positiva Seguros y Reaseguros S.A.A. para que efectúe el acuse de recibo mediante la confirmación de recepción de la notificación remitida por este despacho a su Casilla Electrónica, dentro de los cinco (5) primeros días hábiles siguientes a la fecha en que recibe la notificación.")
    elif mode == "resuelve" and text.startswith("DÉCIMO SEGUNDO:"):
        p._element.getparent().remove(p._element)

doc.save(r'C:\Users\Admin\.gemini\antigravity\scratch\resolucion_temp_r4.docx')
