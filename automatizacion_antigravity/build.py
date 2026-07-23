from pathlib import Path
from typing import Dict

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
import copy
import re

from config import TEMPLATE_PATH
from formato_nucleo import aplicar_formato_final

ANCHORS = {
    "expediente": "0672-2026/CC1",
    "denunciante": "MEDALITH",
    "denunciado": "RÍMAC",
    "fecha": "Lima, 24 de junio de 2026",
    "hechos": "HECHOS",
    "hechos_dummy": "El xx de abril",
    "medida_dummy": "La señora Medina",
    "analisis": "DE LA ADMISIÓN A TRÁMITE DE LA DENUNCIA",
    "resuelve": "RESOLUCIÓN DE LA SECRETARÍA TÉCNICA",
}


def clean_run(run):
    run.font.name = "Arial Narrow"
    # Protect initials size
    if run.text and re.match(r"^[A-Z]{2,4}/[A-Z]{2,4}$", run.text.strip()):
        run.font.size = Pt(8)
    else:
        run.font.size = Pt(11)
    run.font.highlight_color = None
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial Narrow")
    rpr = run._element.get_or_add_rPr()
    for tag in ("w:shd", "w:highlight"):
        element = rpr.find(qn(tag))
        if element is not None:
            rpr.remove(element)


def clean_paragraph_shading(paragraph):
    ppr = paragraph._element.get_or_add_pPr()
    for tag in ("w:shd", "w:highlight"):
        element = ppr.find(qn(tag))
        if element is not None:
            ppr.remove(element)


def _strip_manual_enumerator(text):
    return re.sub(r"^(?:\([ivxIVX]+\)|[0-9]+[\.\)])\s*", "", text)


def _clone_paragraph_before(paragraph):
    new_p_element = copy.deepcopy(paragraph._element)
    paragraph._element.addprevious(new_p_element)
    return Paragraph(new_p_element, paragraph._parent)


def remove_numbering(paragraph):
    numpr = paragraph._element.get_or_add_pPr().find(qn("w:numPr"))
    if numpr is not None:
        paragraph._element.pPr.remove(numpr)


def add_run(paragraph, text, bold=False):
    run = paragraph.add_run(text)
    run.bold = bold
    clean_run(run)
    return run


def _replace(paragraph, text, bold=False):
    paragraph.text = ""
    add_run(paragraph, text, bold)


def _remove(paragraph):
    parent = paragraph._element.getparent()
    if parent is not None:
        parent.remove(paragraph._element)

def _remove_next_empty_p(paragraph):
    """Elimina el siguiente párrafo si está vacío, ignorando bookmarks intermedios."""
    from docx.text.paragraph import Paragraph
    nxt = paragraph._element.getnext()
    while nxt is not None and not nxt.tag.endswith('}p'):
        nxt = nxt.getnext()
    if nxt is not None and nxt.tag.endswith('}p'):
        nxt_p = Paragraph(nxt, paragraph._parent)
        if not nxt_p.text.strip():
            _remove(nxt_p)


def _set_list_paragraph(paragraph, text):
    paragraph.text = ""
    add_run(paragraph, _strip_manual_enumerator(text))

def format_resolutiva_name(name):
    match = re.match(r'^(.*?)\s*\((.*?)\)$', name.strip())
    if match:
        name = match.group(1)
    name = name.replace('RA?MAC', 'RÍMAC').replace('RA-mac', 'Rímac')
    words = name.split()
    for i, w in enumerate(words):
        if w.upper() in ['S.A.', 'S.A.C.', 'E.I.R.L.', 'S.R.L.', 'S.A.A.']:
            words[i] = w.upper()
        elif w.upper() in ['Y', 'E', 'DE', 'DEL', 'EL', 'LA', 'LOS', 'LAS']:
            if i > 0:
                words[i] = w.lower()
            else:
                words[i] = w.capitalize()
        else:
            words[i] = w.capitalize()
    return ' '.join(words)

def preprocess_caso(caso):
    abbreviation = "Rímac"
    if 'denunciado' in caso:
        if 'RIMAC' in caso['denunciado'].upper() or 'RÍMAC' in caso['denunciado'].upper() or 'RA?MAC' in caso['denunciado'].upper():
            caso['denunciado'] = re.sub(r'\([^\)]*R[IÍA\?\-]+MAC SEGUROS[^\)]*\)', '(RÍMAC)', caso['denunciado'], flags=re.IGNORECASE)
            caso['denunciado'] = caso['denunciado'].replace('RA?MAC', 'RÍMAC').replace('RA-mac', 'Rímac')
        
        match = re.match(r'^(.*?)\s*\((.*?)\)$', caso['denunciado'].strip())
        if match:
            abbrev_raw = match.group(2)
            if abbrev_raw.upper() == 'RÍMAC' or abbrev_raw.upper() == 'RIMAC' or abbrev_raw.upper() == 'RA?MAC':
                abbreviation = "Rímac"
            else:
                abbreviation = abbrev_raw.title()
    
    for key, value in caso.items():
        if isinstance(value, str):
            value = value.replace('Rímac Seguros', 'Rímac').replace('Rimac Seguros', 'Rimac').replace('RA-mac Seguros', 'Rímac').replace('RÍMAC SEGUROS', 'RÍMAC')
            value = re.sub(r'(?i)\b(Ley|Decreto Legislativo)\s+N(?:[^\d\s]+)?\s+', r'\1 ', value)
            if key == "medida_correctiva":
                value = re.sub(r'(?i)\bla aseguradora\b', abbreviation, value)
                value = re.sub(r'(?i)\bla compañ[ií]a aseguradora\b', abbreviation, value)
            else:
                value = re.sub(r'(?i)\bla aseguradora\b', lambda m: 'La compañía aseguradora' if m.group(0)[0].isupper() else 'la compañía aseguradora', value)
            if key == "req_info" and "presentar todas las comunicaciones cursadas" not in value.lower():
                roman_numerals = ['i', 'ii', 'iii', 'iv', 'v', 'vi', 'vii', 'viii', 'ix', 'x']
                matches = list(re.finditer(r'\(([ivx]+)\)', value))
                if matches:
                    last_match = matches[-1]
                    last_roman = last_match.group(1).lower()
                    try:
                        next_roman = roman_numerals[roman_numerals.index(last_roman) + 1]
                        pattern = r'(?i)(?:;\s*)?y,?\s*\(' + last_roman + r'\)'
                        if re.search(pattern, value):
                            value = re.sub(pattern, f'; ({last_roman})', value)
                        value = value.rstrip('. ')
                        value += f'; y, ({next_roman}) presentar todas las comunicaciones cursadas con la parte denunciante en virtud de los hechos materia de denuncia.'
                    except ValueError:
                        pass
            caso[key] = value
        elif isinstance(value, list):
            new_list = []
            for item in value:
                if isinstance(item, str):
                    item = item.replace('Rímac Seguros', 'Rímac').replace('Rimac Seguros', 'Rimac').replace('RA-mac Seguros', 'Rímac').replace('RÍMAC SEGUROS', 'RÍMAC')
                    item = re.sub(r'(?i)\b(Ley|Decreto Legislativo)\s+N(?:[^\d\s]+)?\s+', r'\1 ', item)
                    if key == "medida_correctiva":
                        item = re.sub(r'(?i)\bla aseguradora\b', abbreviation, item)
                        item = re.sub(r'(?i)\bla compañ[ií]a aseguradora\b', abbreviation, item)
                    else:
                        item = re.sub(r'(?i)\bla aseguradora\b', lambda m: 'La compañía aseguradora' if m.group(0)[0].isupper() else 'la compañía aseguradora', item)
                new_list.append(item)
            caso[key] = new_list
        elif isinstance(value, dict):
            new_dict = {}
            for k, v in value.items():
                if isinstance(v, str):
                    v = v.replace('Rímac Seguros', 'Rímac').replace('Rimac Seguros', 'Rimac').replace('RA-mac Seguros', 'Rímac').replace('RÍMAC SEGUROS', 'RÍMAC')
                    v = re.sub(r'(?i)\b(Ley|Decreto Legislativo)\s+N(?:[^\d\s]+)?\s+', r'\1 ', v)
                new_dict[k] = v
            caso[key] = new_dict
    return caso

def _default_resolutivos(caso):
    denunciado = format_resolutiva_name(caso["denunciado"])
    denunciante = format_resolutiva_name(caso["denunciante"])
    fecha = caso["fecha_denuncia"]
    return {
        "PRIMERO": f"admitir a trámite la denuncia del {fecha} interpuesta por {denunciante} contra {denunciado}, por lo siguiente:",
        "SEGUNDO": f"tener por ofrecidos los medios probatorios presentados en el escrito de denuncia del {fecha}.",
        "TERCERO": f"requerir a {denunciado} que cumpla con lo siguiente:",
        "CUARTO": f"correr traslado de la denuncia interpuesta el {fecha} a {denunciado}, para que, de conformidad con lo dispuesto por el artículo 26 de la Ley sobre Facultades, Normas y Organización del Indecopi, aprobada por Decreto Legislativo 807, presente sus descargos en un plazo no mayor de cinco (5) días hábiles contados desde la notificación.",
        "QUINTO": f"{caso.get('req_info', '')}",
    }


def build(caso: Dict, template_path: Path = TEMPLATE_PATH, output_path: Path = None) -> Path:
    caso = preprocess_caso(caso)
    template_path = Path(template_path)
    if not template_path.exists():
        raise FileNotFoundError(f"No existe la plantilla: {template_path}")
    doc = Document(template_path)
    for paragraph in doc.paragraphs:
        clean_paragraph_shading(paragraph)
        for run in paragraph.runs:
            clean_run(run)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    clean_paragraph_shading(paragraph)
                    for run in paragraph.runs:
                        clean_run(run)

    found = set()
    analysis_index = resolution_index = notif_index = 0
    mode = "header"
    resolution_text = _default_resolutivos(caso)
    resolution_text.update(caso.get("resolutivos", {}))
    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        if "EXPEDIENTE" in text and ANCHORS["expediente"] in text:
            _replace(paragraph, f"EXPEDIENTE\t:\t{caso['expediente']}", True)
            found.add("EXPEDIENTE")
        elif "DENUNCIANTE" in text and ANCHORS["denunciante"] in text:
            _replace(paragraph, f"DENUNCIANTE\t:\t{caso['denunciante'].upper()}", True)
            found.add("DENUNCIANTE")
        elif "DENUNCIADO" in text and ANCHORS["denunciado"] in text:
            _replace(paragraph, f"DENUNCIADO\t:\t{caso['denunciado'].upper()}", True)
            found.add("DENUNCIADO")
        elif text.startswith("Lima,") and ANCHORS["fecha"] in text:
            _replace(paragraph, caso["fecha_res"])
            found.add("FECHA")
        elif text == ANCHORS["hechos"]:
            mode = "hechos"
            found.add("HECHOS")
        elif mode == "hechos" and text.startswith("Mediante el escrito"):
            _replace(paragraph, f"Mediante el escrito del {caso['fecha_denuncia']}, {caso['intro_denuncia']} señalando lo siguiente:")
            mode = "hechos_list"
            found.add("INTRO_HECHOS")
        elif mode == "hechos_list" and text.startswith(ANCHORS["hechos_dummy"]):
            for i, hecho in enumerate(caso["hechos"]):
                new_paragraph = _clone_paragraph_before(paragraph)
                _set_list_paragraph(new_paragraph, hecho)
                # Only add empty paragraph if it's NOT the last item
                # The last item will use the template's empty paragraph as spacing
                if i < len(caso["hechos"]) - 1:
                    from docx.oxml import OxmlElement
                    empty_p = OxmlElement('w:p')
                    paragraph._element.addprevious(empty_p)
            _remove(paragraph)
            found.add("BLOQUE_HECHOS")
        elif mode == "hechos_list" and text.startswith("El 5 de abril"):
            _remove_next_empty_p(paragraph)
            _remove(paragraph)
        elif mode == "hechos_list" and text.startswith(ANCHORS["medida_dummy"]):
            _set_list_paragraph(paragraph, caso["medida_correctiva"])
            found.add("MEDIDA_CORRECTIVA")
        elif text == ANCHORS["analisis"]:
            mode = "analisis"
            found.add("ANALISIS")
        elif mode == "analisis" and (text.startswith("La Secretaría Técnica") or text.startswith("Así también") or text.startswith("Asimismo") or text.startswith("Además")):
            if analysis_index < len(caso["imputaciones_analisis"]):
                _replace(paragraph, caso["imputaciones_analisis"][analysis_index])
                analysis_index += 1
                found.add("BLOQUE_ANALISIS")
            else:
                # Remove following empty paragraph for discarded items
                _remove_next_empty_p(paragraph)
                _remove(paragraph)
        elif mode == "analisis" and text.startswith("En tanto la denuncia"):
            mode = "req"
        elif mode == "req" and text.startswith("A efectos de tener mayores elementos"):
            _replace(paragraph, caso["req_info"])
            found.add("REQ_INFO")
        elif text == ANCHORS["resuelve"]:
            mode = "resuelve"
            found.add("RESUELVE")
        elif mode == "resuelve" and text.startswith("PRIMERO"):
            _replace(paragraph, "PRIMERO: ", True)
            add_run(paragraph, resolution_text["PRIMERO"])
            found.add("PRIMERO")
        elif mode == "resuelve" and text.startswith("DÉCIMO"):
            label = text.split(":", 1)[0]
            notificaciones = caso.get("notificaciones", [])
            if notif_index < len(notificaciones):
                _replace(paragraph, f"{label}: ", True)
                add_run(paragraph, notificaciones[notif_index])
                notif_index += 1
            else:
                _remove(paragraph)
        elif mode == "resuelve" and text.startswith("Presunta infracción"):
            import copy
            if "global_list_p_xml" not in caso:
                caso["global_list_p_xml"] = copy.deepcopy(paragraph._element)
            if "BLOQUE_RES" not in found:
                for imp in caso["imputaciones_res"]:
                    new_p = _clone_paragraph_before(paragraph)
                    _set_list_paragraph(new_p, imp)
                    from docx.oxml import OxmlElement
                    empty_p = OxmlElement('w:p')
                    paragraph._element.addprevious(empty_p)
                resolution_index = len(caso["imputaciones_res"])
                found.add("BLOQUE_RES")
            
            # Remove any immediately following empty paragraph
            _remove_next_empty_p(paragraph)
            
            _remove(paragraph)
        elif mode == "resuelve" and text.startswith(("SEGUNDO:", "TERCERO:", "CUARTO:", "QUINTO:", "SEXTO:", "SÉTIMO:", "SÉPTIMO:")):
            roman = text.split(":", 1)[0]
            if roman not in resolution_text:
                found.add(roman)
                continue
            
            _replace(paragraph, f"{roman}: ", True)
            
            # The text might be multi-line in JSON
            parts = resolution_text.get(roman, "").split("\\n")
            add_run(paragraph, parts[0])
            found.add(roman)
            
            # If there are sub-items, insert them as new paragraphs
            if len(parts) > 1:
                # We need a template paragraph for lists. We will clone paragraph and change its style.
                # Since we stored global_list_p_xml from BLOQUE_RES, we can insert it.
                from docx.oxml import OxmlElement
                import copy
                from docx.text.paragraph import Paragraph
                
                # Insert parts in reverse order or just move the paragraph pointer
                for req in parts[1:]:
                    if "global_list_p_xml" in caso:
                        new_xml = copy.deepcopy(caso["global_list_p_xml"])
                        paragraph._element.addnext(new_xml)
                        new_p = Paragraph(new_xml, paragraph._parent)
                        _set_list_paragraph(new_p, req)
                        # move the paragraph pointer so they stay in order
                        paragraph = new_p
                    else:
                        # Fallback if no template is found
                        new_p = _clone_paragraph_before(paragraph)
                        _set_list_paragraph(new_p, req)
                        paragraph._element.addnext(new_p._element)
                        paragraph = new_p
        elif text.startswith("Firmado digitalmente por"):
            from docx.text.paragraph import Paragraph
            prev = paragraph._element.getprevious()
            while prev is not None:
                next_prev = prev.getprevious()
                if prev.tag.endswith('}p'):
                    prev_p = Paragraph(prev, paragraph._parent)
                    if not prev_p.text.strip():
                        _remove(prev_p)
                    else:
                        break
                prev = next_prev
            
            from docx.oxml import OxmlElement
            for _ in range(3):
                empty_p = OxmlElement('w:p')
                paragraph._element.addprevious(empty_p)

    missing = {"EXPEDIENTE", "DENUNCIANTE", "DENUNCIADO", "FECHA", "HECHOS", "BLOQUE_HECHOS",
               "ANALISIS", "BLOQUE_ANALISIS", "RESUELVE", "PRIMERO", "SEGUNDO", "TERCERO",
               "CUARTO", "QUINTO"} - found
    if analysis_index != len(caso["imputaciones_analisis"]):
        missing.add(f"ANALISIS ({analysis_index}/{len(caso['imputaciones_analisis'])})")
    if resolution_index != len(caso["imputaciones_res"]):
        missing.add(f"IMPUTACIONES_RES ({resolution_index}/{len(caso['imputaciones_res'])})")
    if "notificaciones" in caso and notif_index != len(caso["notificaciones"]):
        missing.add(f"NOTIFICACIONES ({notif_index}/{len(caso['notificaciones'])})")
    if missing:
        raise RuntimeError("No se pudieron reemplazar las secciones obligatorias: " + ", ".join(sorted(missing)))
    if output_path is None:
        raise ValueError("output_path es obligatorio")
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    aplicar_formato_final(doc)
    doc.save(output_path)
    return output_path
