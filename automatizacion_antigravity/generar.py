import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from build import build
from config import PRODUCTOS_DIR
from insert_footnotes import insert_footnotes
from docx import Document


def _remove_footnote_markers(path):
    doc = Document(path)
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.text = re.sub(r"__F\d*__", "", run.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.text = re.sub(r"__F\d*__", "", run.text)
    doc.save(path)


def main(argv=None):
    argv = argv or sys.argv[1:]
    if len(argv) != 1:
        print("Uso: python generar.py casos/<archivo>.json", file=sys.stderr)
        return 2
    case_path = Path(argv[0])
    if not case_path.exists():
        print(f"No existe el caso: {case_path}", file=sys.stderr)
        return 2
    caso = json.loads(case_path.read_text(encoding="utf-8"))
    output_path = PRODUCTOS_DIR / f"ADM {caso['expediente'].replace('/', '-')}.docx"
    build(caso, output_path=output_path)
    inserted = insert_footnotes(output_path, output_path, caso.get("footnotes", {}))
    if not inserted:
        _remove_footnote_markers(output_path)
    print(f"Generado: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
