import docx
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
import copy
import sys
import os

def clean_run(run):
    run.font.name = 'Arial Narrow'
    run.font.size = Pt(11)
    run.font.highlight_color = None
    run.font.color.rgb = RGBColor(0, 0, 0)
    rPr = run._element.get_or_add_rPr()
    shd = rPr.find(qn('w:shd'))
    if shd is not None:
        rPr.remove(shd)
    highlight = rPr.find(qn('w:highlight'))
    if highlight is not None:
        rPr.remove(highlight)

def clean_paragraph_shading(p):
    pPr = p._element.get_or_add_pPr()
    shd = pPr.find(qn('w:shd'))
    if shd is not None:
        pPr.remove(shd)
    # also remove paragraph level highlight if any (rare but possible)
    highlight = pPr.find(qn('w:highlight'))
    if highlight is not None:
        pPr.remove(highlight)

def remove_numbering(p):
    pPr = p._element.get_or_add_pPr()
    numPr = pPr.find(qn('w:numPr'))
    if numPr is not None:
        pPr.remove(numPr)

def add_run(p, text, bold=False):
    run = p.add_run(text)
    run.bold = bold
    clean_run(run)
    return run

doc = docx.Document(r'D:\CC1\Modelos al 30-06-26\MODELO xxxx exp. ADM ver. David 1 DDO,  VARIAS IMPUTACIONES.docx')

expediente = "2190-2026/CC1"
denunciante = "GLADIS YSABEL BENITES CAJA (SEÑORA BENITES)"
denunciado = "RÍMAC SEGUROS Y REASEGUROS S.A. (RÍMAC SEGUROS)"
fecha_res = "Lima, 17 de julio de 2026"
fecha_denuncia = "12 de junio de 2026"

hechos = [
    "(i)\tEl 24 de julio de 2024, el vehículo de placa de rodaje CMJ-003, de propiedad de la señora Benites y conducido por su esposo, sufrió un accidente de tránsito.",
    "(ii)\tEl siniestro fue reportado a Rímac Seguros, asignándosele el siniestro N° 879294; sin embargo, la aseguradora no envió a ningún representante ni ajustador al lugar del accidente para la verificación de las circunstancias.",
    "(iii)\tMediante Carta N° SVCRCSL-877260 del 05 de agosto de 2024, la aseguradora rechazó la cobertura del siniestro alegando que el conductor no contaba con licencia de conducir vigente y auténtica, sustentándose en una constancia emitida por el Ministerio de Transportes y Comunicaciones.",
    "(iv)\tPese a requerirlo expresamente mediante comunicación del 06 de septiembre de 2024, la aseguradora no remitió a la denunciante copia de la referida constancia utilizada para el rechazo.",
    "(v)\tLa aseguradora no brindó respuesta a las cartas de reconsideración y comunicaciones remitidas por la asegurada."
]

medida_correctiva = "2.\tLa señora Benites solicitó, en calidad de medida correctiva, que la aseguradora cumpla con otorgar la cobertura del Siniestro N° 879294, incluyendo el resarcimiento por daños materiales al vehículo, daños personales a los ocupantes y demás coberturas contratadas, además de imponer sanciones y el pago de costas y costos."

imputaciones_analisis = [
    "La Secretaría Técnica de la Comisión de Protección al Consumidor 1 (en adelante, la Secretaría Técnica), en ejercicio de sus facultades, considera que el hecho denunciado, consistente en que la compañía aseguradora habría rechazado de manera injustificada la cobertura del siniestro ocurrido el 24 de julio de 2024, respecto del vehículo de placa de rodaje CMJ-003, pese a que el conductor contaba con licencia vigente; involucraría una presunta afectación a sus expectativas, quien no habría encontrado una correspondencia entre lo que esperaba recibir de parte del proveedor y lo que realmente recibió. Por consiguiente, corresponde calificar el hecho materia de denuncia como una presunta infracción al deber de idoneidad, tipificado en los artículos 18° y 19° del Código.__F1__",
    "Así también, la Secretaría Técnica, en ejercicio de sus facultades, considera que el hecho denunciado, consistente en que la compañía aseguradora no habría enviado a ningún representante ni ajustador al lugar del accidente tras el reporte del siniestro; involucraría una presunta afectación a sus expectativas, quien no habría encontrado una correspondencia entre lo que esperaba recibir de parte del proveedor y lo que realmente recibió. Por consiguiente, corresponde calificar el hecho materia de denuncia como una presunta infracción al deber de idoneidad, tipificado en los artículos 18° y 19° del Código.",
    "Asimismo, la Secretaría Técnica, en ejercicio de sus facultades, considera que el hecho denunciado, consistente en que la compañía aseguradora no habría remitido copia de la constancia del Ministerio de Transportes y Comunicaciones utilizada para sustentar el rechazo de cobertura, pese a haberla solicitado. Por consiguiente, corresponde calificar el hecho materia de denuncia como una presunta infracción al derecho a la información, tipificado en el artículo 1, numeral 1, literal b) y al artículo 2 del Código.",
    "Además, la Secretaría Técnica, en ejercicio de sus facultades, considera que el hecho denunciado, consistente en que la compañía aseguradora no habría brindado respuesta a las cartas de reconsideración y comunicaciones remitidas por la asegurada. Por consiguiente, corresponde calificar el hecho materia de denuncia como una presunta infracción al derecho a la información, tipificado en el artículo 1, numeral 1, literal b) y al artículo 2 del Código."
]

req_info = "A efectos de tener mayores elementos que sirvan para la resolución definitiva del presente caso, la Secretaría Técnica, en ejercicio de las facultades que la ley le confiere, conviene requerir a la compañía aseguradora que, en un plazo no mayor de cinco (5) días hábiles contado a partir del día siguiente de la notificación de la presente resolución, cumpla con lo siguiente: (i) remitir una copia de la Póliza del Seguro Vehicular N° 2101-1442991; y, (ii) remitir una copia de la constancia del Ministerio de Transportes y Comunicaciones utilizada para rechazar la cobertura del siniestro N° 879294."

imputaciones_res = [
    "(i)\tPresunta infracción a los artículos 18° y 19° de la Ley N° 29571, Código de Protección y Defensa del Consumidor, en tanto la compañía aseguradora habría rechazado de manera injustificada la cobertura del siniestro ocurrido el 24 de julio de 2024, respecto del vehículo de placa de rodaje CMJ-003.",
    "(ii)\tPresunta infracción a los artículos 18° y 19° de la Ley N° 29571, Código de Protección y Defensa del Consumidor, en tanto la compañía aseguradora no habría enviado a ningún representante ni ajustador al lugar del accidente tras el reporte del siniestro.",
    "(iii)\tPresunta infracción al artículo 1, numeral 1, literal b) y al artículo 2 de la Ley N° 29571, Código de Protección y Defensa del Consumidor, en tanto la compañía aseguradora no habría remitido copia de la constancia del Ministerio de Transportes y Comunicaciones utilizada para sustentar el rechazo de cobertura.",
    "(iv)\tPresunta infracción al artículo 1, numeral 1, literal b) y al artículo 2 de la Ley N° 29571, Código de Protección y Defensa del Consumidor, en tanto la compañía aseguradora no habría brindado respuesta a las cartas de reconsideración y comunicaciones remitidas por la asegurada."
]

for p in doc.paragraphs:
    clean_paragraph_shading(p)
    for run in p.runs:
        clean_run(run)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                clean_paragraph_shading(p)
                for run in p.runs:
                    clean_run(run)

mode = "header"
hecho_idx = 0
analisis_idx = 0
resuelve_idx = 0

for p in list(doc.paragraphs):
    text = p.text.strip()
    
    if "EXPEDIENTE" in text and "0672-2026/CC1" in text:
        p.text = ""
        add_run(p, f"EXPEDIENTE\t:\t{expediente}", bold=True)
    elif "DENUNCIANTE" in text and "MEDALITH" in text:
        p.text = ""
        add_run(p, f"DENUNCIANTE\t:\t{denunciante}", bold=True)
    elif "DENUNCIADO" in text and "RÍMAC" in text:
        p.text = ""
        add_run(p, f"DENUNCIADO\t:\t{denunciado}", bold=True)
    elif text.startswith("Lima,"):
        p.text = ""
        add_run(p, fecha_res, bold=False)
        p.paragraph_format.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT
    elif text == "HECHOS":
        mode = "hechos_start"
    elif mode == "hechos_start" and text.startswith("Mediante el escrito"):
        p.text = ""
        add_run(p, f"Mediante el escrito de denuncia del {fecha_denuncia}, la señora Benites denunció a Rímac Seguros por presuntas infracciones a la Ley N° 29571, Código de Protección y Defensa del Consumidor (en adelante, Código), señalando lo siguiente:")
        mode = "hechos_list"
    elif mode == "hechos_list" and (text.startswith("El xx de abril") or text.startswith("El 5 de abril")):
        if text.startswith("El xx de abril"):
            for h in hechos:
                np = p.insert_paragraph_before()
                remove_numbering(np)
                np.paragraph_format.left_indent = Inches(0.5)
                np.paragraph_format.first_line_indent = Inches(-0.5)
                np.paragraph_format.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.JUSTIFY
                add_run(np, h)
        p._element.getparent().remove(p._element)
    elif mode == "hechos_list" and text.startswith("La señora Medina"):
        p.text = ""
        remove_numbering(p)
        p.paragraph_format.left_indent = Inches(0)
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.JUSTIFY
        add_run(p, medida_correctiva)
    elif text == "DE LA ADMISIÓN A TRÁMITE DE LA DENUNCIA":
        mode = "analisis"
    elif mode == "analisis" and (text.startswith("La Secretaría Técnica") or text.startswith("Así también") or text.startswith("Asimismo") or text.startswith("Además")):
        if analisis_idx < len(imputaciones_analisis):
            p.text = ""
            add_run(p, imputaciones_analisis[analisis_idx])
            analisis_idx += 1
        else:
            p._element.getparent().remove(p._element)
    elif mode == "analisis" and text.startswith("En tanto la denuncia"):
        mode = "req"
    elif mode == "req" and text.startswith("A efectos de tener mayores elementos"):
        p.text = ""
        add_run(p, req_info)
    elif text == "RESOLUCIÓN DE LA SECRETARÍA TÉCNICA":
        mode = "resuelve"
    elif mode == "resuelve" and text.startswith("PRIMERO"):
        p.text = ""
        add_run(p, "PRIMERO: ", bold=True)
        add_run(p, f"admitir a trámite la denuncia del {fecha_denuncia} interpuesta por la señora Gladis Ysabel Benites Caja contra Rímac Seguros y Reaseguros S.A., por lo siguiente:")
    elif mode == "resuelve" and text.startswith("Presunta infracción"):
        if resuelve_idx < len(imputaciones_res):
            p.text = ""
            remove_numbering(p)
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.5)
            p.paragraph_format.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.JUSTIFY
            add_run(p, imputaciones_res[resuelve_idx])
            resuelve_idx += 1
        else:
            p._element.getparent().remove(p._element)
    elif mode == "resuelve" and text.startswith("SEGUNDO:"):
        p.text = ""
        add_run(p, "SEGUNDO: ", bold=True)
        add_run(p, f"tener por ofrecidos los medios probatorios presentados en el escrito de denuncia del {fecha_denuncia}.")
    elif mode == "resuelve" and text.startswith("TERCERO:"):
        p.text = ""
        add_run(p, "TERCERO: ", bold=True)
        add_run(p, "requerir a Rímac Seguros y Reaseguros S.A. que cumpla con lo siguiente:")
    elif mode == "resuelve" and text.startswith("CUARTO:"):
        p.text = ""
        add_run(p, "CUARTO: ", bold=True)
        add_run(p, f"correr traslado de la denuncia interpuesta el {fecha_denuncia} a Rímac Seguros y Reaseguros S.A., para que, de conformidad con lo dispuesto por el artículo 26° de la Ley sobre Facultades, Normas y Organización del Indecopi, aprobada por Decreto Legislativo 807, presente sus descargos en un plazo no mayor de cinco (5) días hábiles contados desde la notificación.")
    elif mode == "resuelve" and text.startswith("QUINTO:"):
        p.text = ""
        add_run(p, "QUINTO: ", bold=True)
        add_run(p, req_info.replace("A efectos de tener mayores elementos que sirvan para la resolución definitiva del presente caso, la Secretaría Técnica, en ejercicio de las facultades que la ley le confiere, conviene requerir a la compañía aseguradora que, en un plazo no mayor de cinco (5) días hábiles contado a partir del día siguiente de la notificación de la presente resolución, cumpla con lo siguiente:", "requerir que, en un plazo no mayor de cinco (5) días hábiles contado a partir del día siguiente de la notificación de la presente resolución, el proveedor denunciado cumpla con lo siguiente:"))
    elif mode == "resuelve" and text.startswith("DÉCIMO:"):
        p.text = ""
        add_run(p, "DÉCIMO: ", bold=True)
        add_run(p, "requerir a la señora Gladis Ysabel Benites Caja y a Rímac Seguros y Reaseguros S.A. para que, dentro del plazo de dos (2) días hábiles siguientes a la fecha en que reciban la notificación en su bandeja de correo electrónico, efectúen la confirmación de recepción de la notificación remitida por este despacho a su correo electrónico, de conformidad con el segundo párrafo del numeral 4 del artículo 20 del Texto Único Ordenado de la Ley del Procedimiento Administrativo General, aprobado mediante Decreto Supremo 006-2026-JUS, bajo apercibimiento de rehacer el acto de notificación y notificarles conforme al numeral 1 del artículo 20 del citado cuerpo normativo.")
    elif mode == "resuelve" and text.startswith("DÉCIMO PRIMERO:"):
        p._element.getparent().remove(p._element)
    elif mode == "resuelve" and text.startswith("DÉCIMO SEGUNDO:"):
        p._element.getparent().remove(p._element)

doc.save(r'C:\Users\Admin\.gemini\antigravity\scratch\resolucion_temp_exp2190.docx')
