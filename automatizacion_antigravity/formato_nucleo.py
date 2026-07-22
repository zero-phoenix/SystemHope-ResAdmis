# -*- coding: utf-8 -*-
"""
formato_nucleo.py - NUCLEO COMPARTIDO para todos los generadores CC1.
====================================================================
Single source of truth para constantes y funciones de formato.
TODO script generador DEBE importar desde aqui.

USO:
    from formato_nucleo import (
        add_paragraph, add_run, add_blank, add_meta, add_header,
        ROMAN_INDENT, ROMAN_HANGING, NUM_INDENT, NUM_HANGING,
        LIST_INDENT, LIST_HANGING, META_INDENT, META_HANGING,
        RESOL_LIST_INDENT,
    )

REGLAS: 39 (sangrias), 54 (indentacion metadata), 55 (line spacing),
         56 (fuente Arial Narrow), 57 (tamano 11)
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT


# ====================================================================
# CONSTANTES GLOBALES (Reglas 55, 56, 57, 39, 54)
# ====================================================================

# Fuente y tamano
FUENTE_CUERPO       = "Arial Narrow"
TAMANO_CUERPO       = Pt(11)
FUENTE_NOTAS        = "Arial Narrow"
TAMANO_NOTAS        = Pt(8)
FUENTE_INICIALES    = "Arial Narrow"
TAMANO_INICIALES    = Pt(8)
ALINEACION          = WD_ALIGN_PARAGRAPH.JUSTIFY
SPACE_BEFORE        = Pt(0)
SPACE_AFTER         = Pt(0)
LINE_SPACING        = 1.0
FOOTER_TEXT         = "M-CPC-01/03"

# Sangrias CC1 (Regla 39, 54)
ROMAN_INDENT        = Inches(0.39)
ROMAN_HANGING       = Inches(-0.39)
NUM_INDENT          = Inches(0.39)
NUM_HANGING         = Inches(-0.39)
LIST_INDENT         = Inches(0.79)
LIST_HANGING        = Inches(-0.39)
RESOL_LIST_INDENT   = Inches(0.39)
META_INDENT         = Inches(1.48)
META_HANGING        = Inches(-1.48)

# Secciones resolutivas
ARTICULOS_RESOLUTIVOS = [
    "PRIMERO", "SEGUNDO", "TERCERO", "CUARTO",
    "QUINTO", "SEXTO", "S\u00c9TIMO", "OCTAVO",
    "NOVENO", "D\u00c9CIMO", "D\u00c9CIMO PRIMERO",
    "D\u00c9CIMO SEGUNDO", "D\u00c9CIMO TERCERO",
]


# ====================================================================
# HELPER: add_paragraph
# ====================================================================

def add_paragraph(doc, text, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                  italic=False, underline=False, left_indent=None,
                  hanging_indent=None, font_name=None, font_size=None):
    """
    Anyade un parrafo con formato CC1 estandar.
    Por defecto: Arial Narrow 11, spacing 0, line spacing 1.0.
    """
    if font_name is None:
        font_name = FUENTE_CUERPO
    if font_size is None:
        font_size = TAMANO_CUERPO

    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    if left_indent is not None:
        pf.left_indent = left_indent
    if hanging_indent is not None:
        pf.first_line_indent = hanging_indent
    if left_indent is not None:
        tab_stops = pf.tab_stops
        tab_stops.add_tab_stop(left_indent, WD_TAB_ALIGNMENT.LEFT)
    pf.space_before = SPACE_BEFORE
    pf.space_after = SPACE_AFTER
    pf.line_spacing = LINE_SPACING

    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = font_size
    run.bold = bold
    run.italic = italic
    run.underline = underline
    return p


# ====================================================================
# HELPER: add_run
# ====================================================================

def add_run(p, text, bold=False, italic=False, underline=False,
            font_name=None, font_size=None):
    """
    Anyade un run con formato CC1 estandar.
    Por defecto: Arial Narrow 11.
    """
    if font_name is None:
        font_name = FUENTE_CUERPO
    if font_size is None:
        font_size = TAMANO_CUERPO

    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = font_size
    run.bold = bold
    run.italic = italic
    run.underline = underline
    return run


# ====================================================================
# HELPER: add_blank
# ====================================================================

def add_blank(doc):
    """Anyade un parrafo en blanco con formato CC1."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = SPACE_BEFORE
    pf.space_after = SPACE_AFTER
    pf.line_spacing = LINE_SPACING
    return p


# ====================================================================
# HELPER: add_meta
# ====================================================================

def add_meta(doc, field, value):
    """
    Anyade linea de metadata con sangria colgante de 1.48".
    El field se muestra en bold y value en bold (estandar CC1).
    """
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = META_INDENT
    pf.first_line_indent = META_HANGING
    pf.space_before = SPACE_BEFORE
    pf.space_after = SPACE_AFTER
    pf.line_spacing = LINE_SPACING
    tab_stops = pf.tab_stops
    tab_stops.add_tab_stop(META_INDENT, WD_TAB_ALIGNMENT.LEFT)

    run_f = p.add_run(field)
    run_f.font.name = FUENTE_CUERPO
    run_f.font.size = TAMANO_CUERPO
    run_f.bold = True

    run_v = p.add_run(value)
    run_v.font.name = FUENTE_CUERPO
    run_v.font.size = TAMANO_CUERPO
    run_v.bold = True


# ====================================================================
# HELPER: add_header
# ====================================================================

def add_header(doc, texto=None):
    """
    Anyade el encabezado estandar CC1 (Secretaria Tecnica...).
    Alineacion derecha, bold + italic, Arial Narrow 11.
    """
    if texto is None:
        texto = ("SECRETAR\u00cdA T\u00c9CNICA DE LA\n"
                 "COMISI\u00d3N DE PROTECCI\u00d3N AL CONSUMIDOR 1\n"
                 "SEDE CENTRAL")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pf = p.paragraph_format
    pf.space_before = SPACE_BEFORE
    pf.space_after = SPACE_AFTER
    pf.line_spacing = LINE_SPACING
    run = p.add_run(texto)
    run.font.name = FUENTE_CUERPO
    run.font.size = TAMANO_CUERPO
    run.bold = True
    run.italic = True
    return p

def aplicar_formato_final(doc):
    """
    Aplica las correcciones finales de formato al documento:
    1. Centra la firma digital.
    2. Elimina espaciados innecesarios.
    """
    for p in doc.paragraphs:
        pf = p.paragraph_format
        
        # Eliminar espaciado innecesario si no es necesario (Regla 55)
        # Limitamos space_before a 0 y space_after a 0
        if pf.space_before is not None:
            pf.space_before = Pt(0)
        if pf.space_after is not None and pf.space_after > Pt(0):
            pf.space_after = Pt(0)
            
        # Centrar la firma si existe
        text_upper = p.text.upper()
        if "FIRMADO DIGITALMENTE POR" in text_upper and "EVELING" in text_upper:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
