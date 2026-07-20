import copy
from docx import Document
from docx.text.paragraph import Paragraph

def _clone_paragraph_before(paragraph):
    new_p_element = copy.deepcopy(paragraph._element)
    paragraph._element.addprevious(new_p_element)
    return Paragraph(new_p_element, paragraph._parent)

d = Document(r"Modelos al 30-06-26\MODELO xxxx exp. ADM ver. David 1 DDO,  VARIAS IMPUTACIONES.docx")
for i, p in enumerate(d.paragraphs):
    if "El xx de abril" in p.text:
        print("Found target list paragraph!")
        new_p = _clone_paragraph_before(p)
        new_p.text = ""
        new_p.add_run("Cloned text!")
        
        # verify numPr is there
        numpr = new_p._p.pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr")
        print("Cloned numPr present:", numpr is not None)
        break
