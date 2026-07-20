import re
import sys
from pathlib import Path

from docx import Document


def main(argv=None):
    argv = argv or sys.argv[1:]
    if len(argv) != 1:
        print("Uso: python verificar.py <ruta.docx>")
        return 2
    path = Path(argv[0])
    failures = []
    try:
        doc = Document(path)
    except Exception as exc:
        print(f"FALLA: abrir DOCX: {exc}")
        return 1
    runs = [r for p in doc.paragraphs for r in p.runs if r.text.strip()]
    font_ok = all((r.font.name == "Arial Narrow" or r.font.name is None) and (r.font.size is None or r.font.size.pt == 11) for r in runs)
    print(f"{'OK' if font_ok else 'FALLA'}: cuerpo Arial Narrow 11")
    if not font_ok:
        failures.append("cuerpo no está completamente en Arial Narrow 11")
    footer = "\n".join(p.text for s in doc.sections for p in s.footer.paragraphs)
    footer_ok = "M-CPC-01/03" in footer
    print(f"{'OK' if footer_ok else 'FALLA'}: footer M-CPC-01/03")
    if not footer_ok:
        failures.append("footer no contiene M-CPC-01/03")
    body = "\n".join(p.text for p in doc.paragraphs)
    dummy_literals = (
        "MEDALITH",
        "xx de abril",
        "La señora Medina",
        "0672-2026",
        "__F",
        "{{",
        "Asencios",
        "Mapfre",
        "Comité de Administración del Fondo de Asistencia",
        "Oswaldo Chamorro",
        "poligráfica",
        "1000001429",
        "2101-1031084",
        "30150448",
        "CRT-854049",
        "1364266",
    )
    dummy_patterns = (r"\(\(\(", r"\(N+\)", r"x{3,}", r"\dx{2,}", r"20\dx")
    dummy_matches = [item for item in dummy_literals if item in body]
    dummy_matches.extend(pattern for pattern in dummy_patterns if re.search(pattern, body))
    dummy_ok = not dummy_matches
    print(f"{'OK' if dummy_ok else 'FALLA'}: no quedan textos dummy")
    if not dummy_ok:
        failures.append("quedan textos dummy o marcadores: " + ", ".join(dummy_matches))
    roman_ok = any(
        (p.text.strip().startswith(("I.", "II.", "III.", "IV.", "V."))
         or p.text.strip().startswith(("PRIMERO:", "SEGUNDO:", "TERCERO:", "CUARTO:", "QUINTO:")))
        and any(r.bold for r in p.runs)
        for p in doc.paragraphs
    )
    print(f"{'OK' if roman_ok else 'FALLA'}: títulos romanos en negrita")
    if not roman_ok:
        failures.append("no existen títulos romanos en negrita")
    name_ok = path.name.startswith("ADM ")
    print(f"{'OK' if name_ok else 'FALLA'}: nombre con prefijo ADM")
    if not name_ok:
        failures.append("nombre de archivo no empieza con ADM ")
    if failures:
        print("Fallas:")
        for failure in failures:
            print(f"- {failure}")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
