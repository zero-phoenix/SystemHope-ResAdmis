import sys
from docx import Document

d = Document(r"Modelos al 30-06-26\MODELO xxxx exp. ADM ver. David 1 DDO,  VARIAS IMPUTACIONES.docx")

print("HECHOS LIST PARAGRAPHS (around 'El xx de abril'):")
for i, p in enumerate(d.paragraphs):
    if "El xx de abril" in p.text or "La señora Medina" in p.text:
        print(f"Index {i}: {p.text[:40]}...")
        print("  Style:", p.style.name)
        numpr = p._p.pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr")
        if numpr is not None:
            ilvl = numpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ilvl")
            numId = numpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numId")
            print(f"  numPr: ilvl={ilvl.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') if ilvl is not None else 'None'}, numId={numId.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') if numId is not None else 'None'}")
        else:
            print("  numPr: None")
        print(f"  Indents: left={p.paragraph_format.left_indent}, first_line={p.paragraph_format.first_line_indent}")

print("\nRESUELVE LIST PARAGRAPHS (around 'Presunta infracción'):")
for i, p in enumerate(d.paragraphs):
    if "Presunta infracción" in p.text:
        print(f"Index {i}: {p.text[:40]}...")
        print("  Style:", p.style.name)
        numpr = p._p.pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr")
        if numpr is not None:
            ilvl = numpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ilvl")
            numId = numpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numId")
            print(f"  numPr: ilvl={ilvl.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') if ilvl is not None else 'None'}, numId={numId.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') if numId is not None else 'None'}")
        else:
            print("  numPr: None")
        print(f"  Indents: left={p.paragraph_format.left_indent}, first_line={p.paragraph_format.first_line_indent}")
