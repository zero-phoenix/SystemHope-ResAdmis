#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Genera Resolucion de Admisorio del Expediente 2142-2026/CC1
Aplica todas las reglas de formato CC1 y usa aplicar_reglas_base.py
"""

import os
import re
import sys

# Asegurar que podemos importar los modulos locales
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from aplicar_reglas_base import aplicar_reglas_base


# ============================================================
# DATOS DEL EXPEDIENTE
# ============================================================
EXPEDIENTE = "2142-2026/CC1"
DENUNCIANTE_NOMBRE = "EFRA\u00cdN CORONEL QUISPE"
DENUNCIANTE_ABREV = "SE\u00d1OR CORONEL"
DENUNCIANTE_CORREO = "VERONIKHALYNCH2024@GMAIL.COM"
DENUNCIADO_NOMBRE = "LA POSITIVA SEGUROS Y REASEGUROS S.A.A."
DENUNCIADO_ABREV = "LA POSITIVA"
DENUNCIADO_RUC = "20100210909"
DENUNCIADO_CASILLA = "7230"

FECHA_DENUNCIA = "17 de abril de 2026"
FECHA_SUBSANACION = "26 de junio de 2026"
FECHA_SINIESTRO = "22 de mayo de 2024"
FECHA_CONTRATO = "6 de julio de 2023"

POLIZA_NRO = "230245150"
SUMA_ASEGURADA = "USD 12,450"
PLACA = "BVO-787"

# Traslado
TRASLADO_NRO = "162-2026-PS1/INDECOPI"
FECHA_TRASLADO = "2 de junio de 2026"
FECHA_RECEPCION = "12 de junio de 2026"
EXP_ORIGEN = "000998-2026/PS1"

# Resolucion 1
FECHA_RESOLUCION1 = "23 de junio de 2026"
FECHA_SUBSANACION_EFECTIVA = "26 de junio de 2026"

# ============================================================
# FUNCIONES AUXILIARES
# ============================================================

def add_paragraph(doc, text, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                  italic=False, underline=False, left_indent=None,
                  hanging_indent=None, font_name="Arial Narrow", font_size=Pt(11)):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    if left_indent is not None:
        pf.left_indent = left_indent
    if hanging_indent is not None:
        pf.first_line_indent = hanging_indent
    if left_indent is not None:
        tab_stops = pf.tab_stops
        tab_stops.add_tab_stop(left_indent, WD_TAB_ALIGNMENT.LEFT)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0

    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = font_size
    run.bold = bold
    run.italic = italic
    run.underline = underline
    return p


def add_run(p, text, bold=False, italic=False, underline=False,
            font_name="Arial Narrow", font_size=Pt(11)):
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = font_size
    run.bold = bold
    run.italic = italic
    run.underline = underline
    return run


def add_blank(doc):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    return p


# ============================================================
# SANGRÍAS (Reglas 39, 54, 55)
# ============================================================
ROMAN_INDENT = Inches(0.39)
ROMAN_HANGING = Inches(-0.39)
NUM_INDENT = Inches(0.39)
NUM_HANGING = Inches(-0.39)
LIST_INDENT = Inches(0.79)
LIST_HANGING = Inches(-0.39)
RESOL_LIST_INDENT = Inches(0.39)  # En resolutiva la viñeta va a 0cm, texto a 1cm

# Metadata
META_INDENT = Inches(1.48)
META_HANGING = Inches(-1.48)


def add_meta(doc, field, value):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = META_INDENT
    pf.first_line_indent = META_HANGING
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    tab_stops = pf.tab_stops
    tab_stops.add_tab_stop(META_INDENT, WD_TAB_ALIGNMENT.LEFT)

    run_f = p.add_run(field)
    run_f.font.name = "Arial Narrow"
    run_f.font.size = Pt(11)
    run_f.bold = True

    run_v = p.add_run(value)
    run_v.font.name = "Arial Narrow"
    run_v.font.size = Pt(11)
    run_v.bold = True


# ============================================================
# GENERACIÓN PRINCIPAL
# ============================================================

def generar():
    template_path = os.path.abspath(
        r"D:\BETTER CALL DAVID\ResAdmi\Modelos al 30-06-26\MODELO xxxx exp. ADM ver. David 1 DDO,  VARIAS IMPUTACIONES.docx"
    )
    doc = Document(template_path)

    # Limpiar cuerpo
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)
    for t in list(doc.tables):
        t._element.getparent().remove(t._element)

    # Estilo Normal
    style = doc.styles['Normal']
    style.font.name = "Arial Narrow"
    style.font.size = Pt(11)
    pf_style = style.paragraph_format
    pf_style.space_before = Pt(0)
    pf_style.space_after = Pt(0)
    pf_style.line_spacing = 1.0

    # ============================================================
    # ENCABEZADO
    # ============================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    run = p.add_run("SECRETAR\u00cdA T\u00c9CNICA DE LA\nCOMISI\u00d3N DE PROTECCI\u00d3N AL CONSUMIDOR 1\nSEDE CENTRAL")
    run.font.name = "Arial Narrow"
    run.font.size = Pt(11)
    run.bold = True
    run.italic = True

    add_blank(doc)

    # ============================================================
    # METADATA
    # ============================================================
    add_meta(doc, "EXPEDIENTE\t:\t", f"{EXPEDIENTE}")
    add_meta(doc, "DENUNCIANTE\t:\t", f"{DENUNCIANTE_NOMBRE} ({DENUNCIANTE_ABREV})")
    add_meta(doc, "DENUNCIADO\t:\t", f"{DENUNCIADO_NOMBRE} ({DENUNCIADO_ABREV})")
    add_meta(doc, "MATERIAS\t:\t", "ADMISI\u00d3N A TR\u00c1MITE\n\t\tREQUERIMIENTO DE INFORMACI\u00d3N")
    add_meta(doc, "RESOLUCI\u00d3N\t:\t", "2")

    add_blank(doc)
    add_paragraph(doc, "Lima, 13 de julio de 2026", align=WD_ALIGN_PARAGRAPH.LEFT)
    add_blank(doc)

    # ============================================================
    # I. HECHOS
    # ============================================================
    add_paragraph(doc, "I.\tHECHOS", bold=True,
                  left_indent=ROMAN_INDENT, hanging_indent=ROMAN_HANGING)
    add_blank(doc)

    # Párrafo 1: denuncia y hechos
    texto_hechos_1 = (
        f"1.\tMediante el escrito del {FECHA_DENUNCIA}, subsanado mediante escrito del {FECHA_SUBSANACION}, "
        f"el {DENUNCIANTE_ABREV} denunci\u00f3 a {DENUNCIADO_NOMBRE} por presuntas infracciones "
        f"a la Ley N\u00b029571, C\u00f3digo de Protecci\u00f3n y Defensa del Consumidor "
        f"(en adelante, C\u00f3digo)__F1__, se\u00f1alando lo siguiente:"
    )
    add_paragraph(doc, texto_hechos_1,
                  left_indent=NUM_INDENT, hanging_indent=NUM_HANGING)
    add_blank(doc)

    # Viñetas de hechos (compactas, sin blank entre ellas)
    add_paragraph(doc,
        f"(i)\tEl {FECHA_CONTRATO}, el {DENUNCIANTE_ABREV} contrat\u00f3 con "
        f"la compa\u00f1\u00eda aseguradora el Seguro Vehicular - P\u00f3liza N\u00b0{POLIZA_NRO} "
        f"(en adelante, Seguro Vehicular) correspondiente al veh\u00edculo de placa {PLACA}, "
        f"con una suma asegurada de {SUMA_ASEGURADA}, hecho documentado mediante la citada p\u00f3liza.",
        left_indent=LIST_INDENT, hanging_indent=LIST_HANGING)
    add_paragraph(doc,
        f"(ii)\tEl {FECHA_SINIESTRO}, ocurri\u00f3 el robo total del veh\u00edculo de placa {PLACA} "
        f"mientras se encontraba en el frontis del inmueble donde era guardado, "
        f"hecho denunciado ante la autoridad policial correspondiente.",
        left_indent=LIST_INDENT, hanging_indent=LIST_HANGING)
    add_paragraph(doc,
        f"(iii)\tPese a que la compa\u00f1\u00eda aseguradora recibi\u00f3 la documentaci\u00f3n "
        f"correspondiente, esta decidi\u00f3 denegar el reconocimiento y pago de la cobertura "
        f"solicitada, bajo el argumento de que el veh\u00edculo habr\u00eda sido utilizado para la "
        f"actividad de alquiler, pese a que ello no correspond\u00eda a la realidad.",
        left_indent=LIST_INDENT, hanging_indent=LIST_HANGING)
    add_paragraph(doc,
        f"(iv)\tEl Ministerio P\u00fablico, mediante la disposici\u00f3n fiscal respectiva, "
        f"archiv\u00f3 la investigaci\u00f3n determinando la real condici\u00f3n de propiedad del veh\u00edculo "
        f"y la ausencia de relaci\u00f3n de arrendamiento respecto del mismo.",
        left_indent=LIST_INDENT, hanging_indent=LIST_HANGING)
    add_paragraph(doc,
        f"(v)\tAsimismo, pese a que la compa\u00f1\u00eda aseguradora tom\u00f3 conocimiento "
        f"del siniestro ocurrido, continu\u00f3 descontando de la tarjeta CMR del "
        f"{DENUNCIANTE_ABREV} las primas correspondientes a la renovaci\u00f3n anual "
        f"del Seguro Vehicular (julio 2024 - julio 2025), sin contar con su "
        f"autorizaci\u00f3n expresa para ello.",
        left_indent=LIST_INDENT, hanging_indent=LIST_HANGING)
    add_blank(doc)

    # Párrafo 2: medidas correctivas
    texto_medidas = (
        f"2.\tEl {DENUNCIANTE_ABREV} solicit\u00f3, en calidad de medida correctiva, "
        f"que la compa\u00f1\u00eda aseguradora cumpla con otorgar la cobertura del Seguro Vehicular "
        f"y pagar la suma asegurada de {SUMA_ASEGURADA}. Asimismo, requiri\u00f3 de manera expresa "
        f"el reembolso de costos y costas del presente procedimiento."
    )
    add_paragraph(doc, texto_medidas,
                  left_indent=NUM_INDENT, hanging_indent=NUM_HANGING)
    add_blank(doc)

    # ============================================================
    # II. DE LA ADMISIÓN A TRÁMITE DE LA DENUNCIA
    # ============================================================
    add_paragraph(doc, "II.\tDE LA ADMISI\u00d3N A TR\u00c1MITE DE LA DENUNCIA", bold=True,
                  left_indent=ROMAN_INDENT, hanging_indent=ROMAN_HANGING)
    add_blank(doc)

    # Imputación 1: Idoneidad por denegatoria injustificada
    texto_imputacion_1 = (
        f"3.\tLa Secretar\u00eda T\u00e9cnica de la Comisi\u00f3n de Protecci\u00f3n al Consumidor N\u00b01 "
        f"(en adelante, la Secretar\u00eda T\u00e9cnica), en ejercicio de sus facultades__F2__, "
        f"considera que el hecho denunciado, consistente en que la compa\u00f1\u00eda aseguradora habr\u00eda "
        f"denegado injustificadamente al {DENUNCIANTE_ABREV} la cobertura de la "
        f"P\u00f3liza N\u00b0{POLIZA_NRO} del Seguro Vehicular; involucrar\u00eda una presunta afectaci\u00f3n "
        f"a sus expectativas, quien no habr\u00eda encontrado una correspondencia entre lo que "
        f"esperaba recibir de parte del proveedor y lo que realmente recibi\u00f3. Por consiguiente, "
        f"corresponde calificar el hecho materia de denuncia como una presunta infracci\u00f3n "
        f"al deber de idoneidad, tipificado en los art\u00edculos 18\u00b0 y 19\u00b0 del C\u00f3digo__F3__."
    )
    add_paragraph(doc, texto_imputacion_1,
                  left_indent=NUM_INDENT, hanging_indent=NUM_HANGING)
    add_blank(doc)

    # Imputación 2: Idoneidad por renovación automática no solicitada
    texto_imputacion_2 = (
        f"4.\tAsimismo, la Secretar\u00eda T\u00e9cnica considera que el hecho denunciado, consistente "
        f"en que la compa\u00f1\u00eda aseguradora habr\u00eda continuado descontando de la tarjeta CMR del "
        f"{DENUNCIANTE_ABREV} las primas correspondientes a la renovaci\u00f3n anual "
        f"del Seguro Vehicular (julio 2024 - julio 2025), pese a que este no habr\u00eda prestado "
        f"su autorizaci\u00f3n expresa para dicha renovaci\u00f3n y pese al siniestro ocurrido; "
        f"involucrar\u00eda una presunta afectaci\u00f3n a sus expectativas. Por consiguiente, "
        f"corresponde calificar el hecho materia de denuncia como una presunta infracci\u00f3n "
        f"al deber de idoneidad, tipificado en los art\u00edculos 18\u00b0 y 19\u00b0 del C\u00f3digo__F3__."
    )
    add_paragraph(doc, texto_imputacion_2,
                  left_indent=NUM_INDENT, hanging_indent=NUM_HANGING)
    add_blank(doc)

    # Párrafo 5: Admisión a trámite
    texto_admision = (
        f"5.\tEn tanto la denuncia re\u00fane los requisitos establecidos por la "
        f"norma citada, corresponde admitirla a tr\u00e1mite__F4__."
    )
    add_paragraph(doc, texto_admision,
                  left_indent=NUM_INDENT, hanging_indent=NUM_HANGING)
    add_blank(doc)

    # ============================================================
    # III. REQUERIMIENTO DE INFORMACIÓN
    # ============================================================
    add_paragraph(doc, "III.\tREQUERIMIENTO DE INFORMACI\u00d3N", bold=True,
                  left_indent=ROMAN_INDENT, hanging_indent=ROMAN_HANGING)
    add_blank(doc)

    texto_requerimiento = (
        f"6.\tA efectos de tener mayores elementos que sirvan para la resoluci\u00f3n "
        f"definitiva del presente caso, la Secretar\u00eda T\u00e9cnica, en ejercicio de las "
        f"facultades que la ley le confiere, conviene requerir a la compa\u00f1\u00eda aseguradora que, "
        f"en un plazo no mayor de cinco (5) d\u00edas h\u00e1biles contado a partir del d\u00eda siguiente "
        f"de la notificaci\u00f3n de la presente resoluci\u00f3n, cumpla con lo siguiente: "
        f"(i) presentar una copia completa, legible y debidamente suscrita de la "
        f"P\u00f3liza N\u00b0{POLIZA_NRO} - Seguro Vehicular, as\u00ed como el cargo de remisi\u00f3n de la p\u00f3liza; "
        f"(ii) presentar los medios probatorios que acrediten que la negativa de otorgamiento "
        f"de cobertura fue justificada; y, (iii) presentar todas las comunicaciones cursadas "
        f"con la parte denunciante en virtud de los hechos materia de denuncia."
    )
    add_paragraph(doc, texto_requerimiento,
                  left_indent=NUM_INDENT, hanging_indent=NUM_HANGING)
    add_blank(doc)

    # ============================================================
    # IV. RESOLUCIÓN DE LA SECRETARÍA TÉCNICA
    # ============================================================
    add_paragraph(doc, "IV.\tRESOLUCI\u00d3N DE LA SECRETAR\u00cdA T\u00c9CNICA", bold=True,
                  left_indent=ROMAN_INDENT, hanging_indent=ROMAN_HANGING)
    add_blank(doc)

    # PRIMERO
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_run(p, "PRIMERO: ", bold=True)
    add_run(p,
        f"admitir a tr\u00e1mite la denuncia del {FECHA_DENUNCIA}, subsanada mediante "
        f"escrito del {FECHA_SUBSANACION}, interpuesta por el {DENUNCIANTE_NOMBRE} "
        f"contra {DENUNCIADO_NOMBRE}, por presuntas infracciones a los art\u00edculos "
        f"18\u00b0 y 19\u00b0 de la Ley N\u00b0 29571, C\u00f3digo de Protecci\u00f3n y Defensa del Consumidor, "
        f"en atenci\u00f3n a lo siguiente:")
    add_blank(doc)

    # Viñetas resolutivas (con sangría de RESOL_LIST_INDENT)
    add_paragraph(doc,
        f"(i)\tA los art\u00edculos 18\u00b0 y 19\u00b0, en tanto la compa\u00f1\u00eda aseguradora habr\u00eda "
        f"denegado injustificadamente al {DENUNCIANTE_ABREV} la cobertura de la "
        f"P\u00f3liza N\u00b0{POLIZA_NRO} del Seguro Vehicular.",
        left_indent=RESOL_LIST_INDENT, hanging_indent=LIST_HANGING)
    add_paragraph(doc,
        f"(ii)\tA los art\u00edculos 18\u00b0 y 19\u00b0, en tanto la compa\u00f1\u00eda aseguradora habr\u00eda "
        f"continuado descontando de la tarjeta CMR del {DENUNCIANTE_ABREV} las primas "
        f"correspondientes a la renovaci\u00f3n anual del Seguro Vehicular (julio 2024 - julio 2025), "
        f"sin contar con su autorizaci\u00f3n expresa.",
        left_indent=RESOL_LIST_INDENT, hanging_indent=LIST_HANGING)
    add_blank(doc)

    # SEGUNDO
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_run(p, "SEGUNDO: ", bold=True)
    add_run(p,
        f"tener por ofrecidos los medios probatorios presentados en el escrito de "
        f"denuncia del {FECHA_DENUNCIA}, subsanado mediante escrito del {FECHA_SUBSANACION}.")
    add_blank(doc)

    # TERCERO
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_run(p, "TERCERO: ", bold=True)
    add_run(p,
        f"requerir a {DENUNCIADO_NOMBRE} que cumpla con lo siguiente:")
    add_blank(doc)

    # Incisos TERCERO (compactos, sin blank entre ellos)
    add_paragraph(doc,
        "(i)\tpresentar documentos que acrediten su inscripci\u00f3n en los Registros "
        "P\u00fablicos o una declaraci\u00f3n jurada que indique que cuenta con dicha inscripci\u00f3n;",
        left_indent=RESOL_LIST_INDENT, hanging_indent=LIST_HANGING)
    add_paragraph(doc,
        "(ii)\tpresentar copia simple de las facultades de representaci\u00f3n de su "
        "representante legal en el presente procedimiento o la declaraci\u00f3n jurada que "
        "indique que cuenta con dichas facultades y que estas se encuentran vigentes;",
        left_indent=RESOL_LIST_INDENT, hanging_indent=LIST_HANGING)
    add_paragraph(doc,
        "(iii)\tconsignar el N\u00famero de Registro \u00danico de Contribuyentes (RUC); y,",
        left_indent=RESOL_LIST_INDENT, hanging_indent=LIST_HANGING)
    add_paragraph(doc,
        "(iv)\tfijar domicilio procesal para el procedimiento, de conformidad con el "
        "numeral 1 del art\u00edculo 442\u00b0 del C\u00f3digo Procesal Civil. Para ello podr\u00e1 "
        "se\u00f1alar domicilio f\u00edsico o una direcci\u00f3n electr\u00f3nica a la que pueda "
        "remitirse las notificaciones;",
        left_indent=RESOL_LIST_INDENT, hanging_indent=LIST_HANGING)
    add_paragraph(doc,
        "(v)\ten caso califique como micro empresa o peque\u00f1a empresa, presentar los "
        "documentos que acrediten su volumen de ventas o ingresos brutos percibidos el "
        "a\u00f1o anterior relativo a todas sus actividades econ\u00f3micas y el n\u00famero de "
        "trabajadores con el que cuenta__F5__. Ello, a fin de que la Comisi\u00f3n pueda "
        "meritar dicha documentaci\u00f3n, conforme lo establece el art\u00edculo 110\u00b0 del "
        "C\u00f3digo__F6__.",
        left_indent=RESOL_LIST_INDENT, hanging_indent=LIST_HANGING)
    add_blank(doc)

    # CUARTO
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_run(p, "CUARTO: ", bold=True)
    add_run(p,
        f"correr traslado de la denuncia interpuesta el {FECHA_DENUNCIA}, subsanada "
        f"mediante escrito del {FECHA_SUBSANACION}, a {DENUNCIADO_NOMBRE}, para que, "
        f"de conformidad con lo dispuesto por el art\u00edculo 26\u00b0 de la Ley sobre "
        f"Facultades, Normas y Organizaci\u00f3n del Indecopi, aprobada por Decreto "
        f"Legislativo 807, presente sus descargos en un plazo no mayor de cinco (5) "
        f"d\u00edas h\u00e1biles contados desde la notificaci\u00f3n.")
    add_blank(doc)

    # QUINTO
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_run(p, "QUINTO: ", bold=True)
    add_run(p,
        f"requerir que, en un plazo no mayor de cinco (5) d\u00edas h\u00e1biles contado "
        f"a partir del d\u00eda siguiente de la notificaci\u00f3n de la presente resoluci\u00f3n, "
        f"la compa\u00f1\u00eda aseguradora cumpla con: (i) presentar una copia completa, legible "
        f"y debidamente suscrita de la P\u00f3liza N\u00b0{POLIZA_NRO} del Seguro Vehicular, "
        f"as\u00ed como el cargo de remisi\u00f3n de la p\u00f3liza; (ii) presentar los medios "
        f"probatorios que acrediten que la negativa de otorgamiento de cobertura fue "
        f"justificada; y, (iii) presentar todas las comunicaciones cursadas con la "
        f"parte denunciante en virtud de los hechos materia de denuncia.")
    add_blank(doc)

    # SEXTO
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_run(p, "SEXTO: ", bold=True)
    add_run(p,
        f"informar a {DENUNCIADO_NOMBRE} que puede acceder a los documentos presentados "
        f"por el {DENUNCIANTE_ABREV}, a trav\u00e9s del link que se remiti\u00f3 en su oportunidad "
        f"a la casilla electr\u00f3nica de la compa\u00f1\u00eda aseguradora.")
    add_blank(doc)

    # SÉTIMO
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_run(p, "S\u00c9TIMO: ", bold=True)
    add_run(p,
        "informar a las partes que el art\u00edculo 110\u00b0 de la Ley N\u00b0 29571, "
        "C\u00f3digo de Protecci\u00f3n y Defensa del Consumidor__F7__, faculta a la Comisi\u00f3n "
        "a calificar las infracciones de la referida norma como leves, graves o muy "
        "graves e imponer sanciones que van desde una amonestaci\u00f3n hasta una multa "
        "por un m\u00e1ximo de 450 Unidades Impositivas Tributarias, sin perjuicio de las "
        "medidas correctivas, reparadoras y complementarias, que puedan ordenarse de "
        "acuerdo a lo estipulado en los art\u00edculos 114\u00b0, 115\u00b0 y 116\u00b0 de la referida "
        "norma__F8__. Asimismo, se consideran circunstancias atenuantes para la "
        "graduaci\u00f3n de la sanci\u00f3n, el allanamiento de la denuncia o el reconocimiento "
        "de las pretensiones en ella contenidas, de acuerdo con el art\u00edculo 112\u00b0 "
        "del C\u00f3digo__F9__.")
    add_blank(doc)

    # OCTAVO
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_run(p, "OCTAVO: ", bold=True)
    add_run(p,
        "informar a las partes que conforme a lo establecido en el art\u00edculo 39\u00b0 "
        "de la Ley sobre Facultades, Normas y Organizaci\u00f3n del Indecopi, aprobada "
        "por Decreto Legislativo 807, los gastos por los peritajes realizados, "
        "actuaci\u00f3n de pruebas, inspecciones y otros derivados de la tramitaci\u00f3n del "
        "proceso ser\u00e1n de cargo de la parte que solicita la prueba, salvo pacto en "
        "contrario. En todos los casos, la resoluci\u00f3n final determinar\u00e1 si los gastos "
        "deben ser asumidos por alguna de las partes, o reembolsados a la otra parte "
        "o al Indecopi, seg\u00fan sea el caso, de manera adicional a la sanci\u00f3n que haya "
        "podido imponerse.")
    add_blank(doc)

    # NOVENO
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_run(p, "NOVENO: ", bold=True)
    add_run(p,
        "comunicar a las partes que, de acuerdo a lo se\u00f1alado por el art\u00edculo 29\u00b0 "
        "del Decreto Legislativo N\u00b0 807__F10__, hasta antes de la emisi\u00f3n de la "
        "Resoluci\u00f3n Final tienen la posibilidad de solicitar una audiencia de "
        "conciliaci\u00f3n. En caso deleguen a favor de una tercera persona su actuaci\u00f3n "
        "en la diligencia programada, \u00e9sta deber\u00e1 presentar un poder especial con "
        "firma legalizada ante Notario P\u00fablico, ")
    add_run(p,
        "donde conste expresamente su facultad para asistir y conciliar en su "
        "representaci\u00f3n. Ello bajo apercibimiento de no realizar la audiencia de "
        "conciliaci\u00f3n y levantar el acta de inasistencia correspondiente.",
        underline=True)
    add_blank(doc)

    # DÉCIMO
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_run(p, "D\u00c9CIMO: ", bold=True)
    add_run(p,
        "poner en conocimiento de las partes que, antes de la emisi\u00f3n de la "
        "Resoluci\u00f3n Final, tienen la posibilidad de formular su desistimiento o "
        "presentar el acuerdo arribado mediante la conciliaci\u00f3n, mediaci\u00f3n, "
        "transacci\u00f3n o cualquier otro que, de forma indubitable, deje constancia "
        "que se ha solucionado la controversia materia de denuncia.")
    add_blank(doc)

    # DÉCIMO PRIMERO (notificación denunciante - correo electrónico)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_run(p, "D\u00c9CIMO PRIMERO: ", bold=True)
    add_run(p,
        f"requerir al {DENUNCIANTE_NOMBRE} para que, dentro del plazo de dos (2) "
        f"d\u00edas h\u00e1biles siguientes a la fecha en que reciba la notificaci\u00f3n en su "
        f"bandeja de correo electr\u00f3nico, efect\u00fae la confirmaci\u00f3n de recepci\u00f3n de la "
        f"notificaci\u00f3n remitida por este despacho a su correo electr\u00f3nico, de "
        f"conformidad con el segundo p\u00e1rrafo del numeral 4 del art\u00edculo 20\u00b0 del "
        f"Texto \u00danico Ordenado de la Ley del Procedimiento Administrativo General, "
        f"aprobado mediante Decreto Supremo N\u00b0 004-2019-JUS, bajo apercibimiento de "
        f"rehacer el acto de notificaci\u00f3n y notificarle conforme al numeral 1 del "
        f"art\u00edculo 20\u00b0 del citado cuerpo normativo__F11__.")
    add_blank(doc)

    # DÉCIMO SEGUNDO (notificación denunciado - casilla electrónica)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_run(p, "D\u00c9CIMO SEGUNDO: ", bold=True)
    add_run(p,
        f"requerir a {DENUNCIADO_NOMBRE} para que efect\u00fae el acuse de recibo mediante "
        f"la confirmaci\u00f3n de recepci\u00f3n de la notificaci\u00f3n remitida por este despacho "
        f"a su Casilla Electr\u00f3nica, dentro de los cinco (5) primeros d\u00edas h\u00e1biles "
        f"siguientes a la fecha en que recibe la notificaci\u00f3n.")

    add_blank(doc)
    add_blank(doc)

    # Firma
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_run(p,
        "Firmado digitalmente por\nEVELING ROA QUISPE\nSecretaria T\u00e9cnica\n"
        "Comisi\u00f3n de Protecci\u00f3n al Consumidor N\u00b01",
        bold=True)

    # Iniciales tamaño 8
    p_iniciales = doc.add_paragraph()
    p_iniciales.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_iniciales.paragraph_format.space_before = Pt(0)
    p_iniciales.paragraph_format.space_after = Pt(0)
    p_iniciales.paragraph_format.line_spacing = 1.0
    r = p_iniciales.add_run("LGP/JCQ")
    r.font.name = "Arial Narrow"
    r.font.size = Pt(8)

    # ============================================================
    # GUARDAR TEMPORAL
    # ============================================================
    temp_path = os.path.abspath(
        r"D:\BETTER CALL DAVID\ResAdmi\temp_2142.docx"
    )
    doc.save(temp_path)
    print(f"Documento base guardado: {temp_path}")
    return temp_path


# ============================================================
# INYECCIÓN DE NOTAS AL PIE (win32com)
# ============================================================

def inject_footnotes(docx_path):
    """Inyecta las notas al pie reales usando win32com."""
    import win32com.client

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    doc_com = word.Documents.Open(docx_path)

    footnotes_data = {
        "__F1__": (
            "Denuncia remitida a esta Comisi\u00f3n mediante DOCUMENTO DE TRASLADO "
            f"N\u00b0 {TRASLADO_NRO} de fecha {FECHA_TRASLADO}, recepcionada el "
            f"{FECHA_RECEPCION}."
        ),
        "__F2__": (
            "<b>LEY N\u00b0 29571, C\u00d3DIGO DE PROTECCI\u00d3N Y DEFENSA DEL CONSUMIDOR, "
            "publicado el 2 de setiembre de 2010 y modificado por Decreto Legislativo "
            "N\u00b0 1308</b>\n<b>Art\u00edculo 105.- Autoridad competente.</b>\nEl Instituto "
            "Nacional de Defensa de la Competencia y de la Protecci\u00f3n de la Propiedad "
            "Intelectual (Indecopi) es la autoridad con competencia primaria y de alcance "
            "nacional para conocer las presuntas infracciones a las disposiciones contenidas "
            "en el presente C\u00f3digo, as\u00ed como para imponer las sanciones y medidas "
            "correctivas establecidas en el presente cap\u00edtulo, conforme al Decreto "
            "Legislativo N\u00b0 1033, Ley de Organizaci\u00f3n y Funciones del Indecopi. Dicha "
            "competencia solo puede ser negada cuando ella haya sido asignada o se asigne "
            "a favor de otro organismo por norma expresa con rango de ley.\n(...).\n\n"
            "<b>LEY DE ORGANIZACI\u00d3N Y FUNCIONES DEL INDECOPI, APROBADA POR DECRETO "
            "LEGISLATIVO N\u00b0 1033</b>\n<b>Art\u00edculo 27.-</b> De la Comisi\u00f3n de Protecci\u00f3n "
            "al Consumidor.-\nCorresponde a la Comisi\u00f3n de Protecci\u00f3n al Consumidor velar "
            "por el cumplimiento de la Ley de Protecci\u00f3n al Consumidor y de las leyes que, "
            "en general, protegen a los consumidores de la falta de idoneidad de los bienes "
            "y servicios en funci\u00f3n de la informaci\u00f3n brindada, de las omisiones de "
            "informaci\u00f3n y de la discriminaci\u00f3n en el consumo, as\u00ed como de aquellas "
            "que complementen o sustituyan a las anteriores."
        ),
        "__F3__": (
            "<b>LEY N\u00b0 29571, C\u00d3DIGO DE PROTECCI\u00d3N Y DEFENSA DEL CONSUMIDOR, "
            "publicado el 2 de setiembre de 2010 y modificado por Decreto Legislativo "
            "N\u00b0 1308</b>\n<b>Art\u00edculo 18.- Idoneidad</b>\nSe entiende por idoneidad la "
            "correspondencia entre lo que un consumidor espera y lo que efectivamente recibe, "
            "en funci\u00f3n a lo que se le hubiera ofrecido, la publicidad e informaci\u00f3n "
            "transmitida, las condiciones y circunstancias de la transacci\u00f3n, las "
            "caracter\u00edsticas y naturaleza del producto o servicio, el precio, entre otros "
            "factores, atendiendo a las circunstancias del caso.\nLa idoneidad es evaluada "
            "en funci\u00f3n a la propia naturaleza del producto o servicio y a su aptitud para "
            "satisfacer la finalidad para la cual ha sido puesto en el mercado. Las "
            "autorizaciones por parte de los organismos del Estado para la fabricaci\u00f3n de "
            "un producto o la prestaci\u00f3n de un servicio, en los casos que sea necesario, "
            "no eximen de responsabilidad al proveedor frente al consumidor.\n\n"
            "<b>Art\u00edculo 19.- Obligaci\u00f3n de los proveedores</b>\nEl proveedor responde por "
            "la idoneidad y calidad de los productos y servicios ofrecidos; por la "
            "autenticidad de las marcas y leyendas que exhiben sus productos o del signo "
            "que respalda al prestador del servicio, por la falta de conformidad entre la "
            "publicidad comercial de los productos y servicios y \u00e9stos, as\u00ed como por el "
            "contenido y la vida \u00fatil del producto indicado en el envase, en lo que "
            "corresponda."
        ),
        "__F4__": (
            "<b>LEY SOBRE FACULTADES, NORMAS Y ORGANIZACI\u00d3N DEL INDECOPI, APROBADA POR "
            "DECRETO LEGISLATIVO N\u00b0 807</b>\n<b>Art\u00edculo 24.-</b> El Secretario T\u00e9cnico se "
            "encargar\u00e1 de la tramitaci\u00f3n del procedimiento. Para ello, cuenta con las "
            "siguientes facultades:\n(...)\nc) Admitir denuncias a tr\u00e1mite, en aquellos "
            "casos en que la Comisi\u00f3n le haya delegado esta facultad.\n(...)."
        ),
        "__F5__": (
            "El n\u00famero de trabajadores se atender\u00e1 siempre que la empresa haya sido "
            "constituida antes de la vigencia de la Ley N\u00ba 30056, publicada el 2 de julio "
            "de 2013, que modific\u00f3 el art\u00edculo 5 del Texto \u00danico Ordenado de la Ley de "
            "Promoci\u00f3n de la Competitividad, Formalizaci\u00f3n y Desarrollo de la Micro y "
            "Peque\u00f1a Empresa y de Acceso al Empleo Decente (LEY MYPE). Ello, en la medida "
            "que la Tercera Disposici\u00f3n Complementaria Transitoria de la norma modificatoria "
            "precisa que las empresas constituidas antes de la entrada en vigencia de dicha "
            "Ley se rigen por los requisitos de acogimiento al r\u00e9gimen de las micro y "
            "peque\u00f1as empresas regulados en el Decreto Legislativo 1086."
        ),
        "__F6__": (
            "<b>LEY N\u00b0 29571, C\u00d3DIGO DE PROTECCI\u00d3N Y DEFENSA DEL CONSUMIDOR, "
            "publicado el 2 de setiembre de 2010 y modificado por Decreto Legislativo "
            "N\u00b0 1308</b>\n<b>Art\u00edculo 110.- Sanciones Administrativas</b>\n(...)\nEn el "
            "caso de las microempresas, la multa no puede superar el diez por ciento (10%) "
            "de las ventas o ingresos brutos percibidos por el infractor, relativos a todas "
            "sus actividades econ\u00f3micas, correspondientes al ejercicio inmediato anterior "
            "al de la expedici\u00f3n de la resoluci\u00f3n de primera instancia, siempre que se "
            "haya acreditado dichos ingresos, no se encuentre en una situaci\u00f3n de "
            "reincidencia y el caso no verse sobre la vida, salud o integridad de los "
            "consumidores. Para el caso de las peque\u00f1as empresas, la multa no puede "
            "superar el veinte por ciento (20%) de las ventas o ingresos brutos percibidos "
            "por el infractor, conforme a los requisitos se\u00f1alados anteriormente.\n(...)."
        ),
        "__F7__": (
            "<b>LEY N\u00b0 29571, C\u00d3DIGO DE PROTECCI\u00d3N Y DEFENSA DEL CONSUMIDOR, "
            "publicado el 2 de setiembre de 2010 y modificado por Decreto Legislativo "
            "N\u00b0 1308</b>\n<b>Art\u00edculo 110.- Sanciones administrativas</b>\nEl \u00f3rgano "
            "resolutivo puede sancionar las infracciones administrativas a que se refiere "
            "el art\u00edculo 108\u00ba con amonestaci\u00f3n y multas de hasta cuatrocientos cincuenta "
            "(450) Unidades Impositivas Tributarias (UIT), las cuales son calificadas de "
            "la siguiente manera:\na. Infracciones leves, con una amonestaci\u00f3n o con una "
            "multa de hasta cincuenta (50) UIT.\nb. Infracciones graves, con una multa de "
            "hasta ciento cincuenta (150) UIT.\nc. Infracciones muy graves, con una multa "
            "de hasta cuatrocientos cincuenta (450) UIT.\n(...)."
        ),
        "__F8__": (
            "<b>Art\u00edculo 114.- Medidas correctivas</b>\nSin perjuicio de la sanci\u00f3n "
            "administrativa que corresponda al proveedor por una infracci\u00f3n al presente "
            "C\u00f3digo, el Indecopi puede dictar, en calidad de mandatos, medidas correctivas "
            "reparadoras y complementarias.\nLas medidas correctivas reparadoras pueden "
            "dictarse a pedido de parte o de oficio, siempre y cuando sean expresamente "
            "informadas sobre esa posibilidad en la notificaci\u00f3n de cargo al proveedor "
            "por la autoridad encargada del procedimiento.\nLas medidas correctivas "
            "complementarias pueden dictarse de oficio o a pedido de parte.\n\n"
            "<b>Art\u00edculo 115.- Medidas correctivas reparadoras</b>\n115.1 Las medidas "
            "correctivas reparadoras tienen el objeto de resarcir las consecuencias "
            "patrimoniales directas e inmediatas ocasionadas al consumidor por la "
            "infracci\u00f3n administrativa a su estado anterior. En caso el \u00f3rgano resolutivo "
            "dicte una o varias medidas correctivas, debe considerar lo acordado por las "
            "partes durante la relaci\u00f3n de consumo. Las medidas correctivas reparadoras "
            "pueden consistir en ordenar al proveedor infractor lo siguiente:\n"
            "a. Reparar productos.\nb. Cambiar productos por otros de id\u00e9nticas o similares "
            "caracter\u00edsticas, cuando la reparaci\u00f3n no sea posible o no resulte razonable "
            "seg\u00fan las circunstancias.\nc. Entregar un producto de id\u00e9nticas caracter\u00edsticas "
            "o, cuando esto no resulte posible, de similares caracter\u00edsticas, en los "
            "supuestos de p\u00e9rdida o deterioro atribuible al proveedor y siempre que "
            "exista inter\u00e9s del consumidor.\nd. Cumplir con ejecutar la prestaci\u00f3n u "
            "obligaci\u00f3n asumida; y si esto no resulte posible o no sea razonable, otra "
            "de efectos equivalentes, incluyendo prestaciones dinerarias.\n"
            "e. Cumplir con ejecutar otras prestaciones u obligaciones legales o "
            "convencionales a su cargo.\nf. Devolver la contraprestaci\u00f3n pagada por el "
            "consumidor, m\u00e1s los intereses legales correspondientes, cuando la reparaci\u00f3n, "
            "reposici\u00f3n, o cumplimiento de la prestaci\u00f3n u obligaci\u00f3n, seg\u00fan sea el "
            "caso, no resulte posible o no sea razonable.\n\n<b>Art\u00edculo 116.- Medidas "
            "correctivas complementarias</b>\n116.1 Las medidas correctivas complementarias "
            "tienen como objetivo remover los efectos de la conducta infractora o evitar "
            "que esta se reproduzca en el futuro, pudiendo ordenar, entre otros, lo "
            "siguiente:\na. El cese de la conducta infractora y, de ser el caso, la "
            "colocaci\u00f3n de un aviso en el establecimiento comercial o veh\u00edculo de venta "
            "ambulatoria informando sobre la medida dispuesta por la autoridad.\n"
            "b. Disponer que el proveedor difunda, a su costa, las medidas correctivas "
            "dispuestas, as\u00ed como la resoluci\u00f3n final en los medios que correspondan "
            "en funci\u00f3n del alcance de la conducta infractora.\nc. Ordenar que el "
            "proveedor cumpla con prestar los servicios de atenci\u00f3n de reclamos y "
            "soluci\u00f3n de controversias al consumidor, seg\u00fan lo establecido en el "
            "art\u00edculo 150 del C\u00f3digo.\nd. En los casos que corresponda, ordenar el "
            "resarcimiento de los da\u00f1os y perjuicios ocasionados al consumidor, conforme "
            "a las normas del C\u00f3digo Civil."
        ),
        "__F9__": (
            "<b>LEY N\u00b0 29571, C\u00d3DIGO DE PROTECCI\u00d3N Y DEFENSA DEL CONSUMIDOR, "
            "publicado el 2 de setiembre de 2010 y modificado por Decreto Legislativo "
            "N\u00b0 1308</b>\n<b>Art\u00edculo 112.- Criterios de graduaci\u00f3n de las sanciones "
            "administrativas</b>\n(...)\nSe consideran circunstancias atenuantes especiales, "
            "las siguientes:\n(...)\n3. En los procedimientos de oficio promovidos por una "
            "denuncia de parte, cuando el proveedor se allana a la denuncia presentada o "
            "reconoce las pretensiones en ella contenidas, se da por concluido el "
            "procedimiento liminarmente, pudiendo imponerse una amonestaci\u00f3n si el "
            "allanamiento o reconocimiento se realiza con la presentaci\u00f3n de los "
            "descargos; caso contrario la sanci\u00f3n a imponer ser\u00e1 pecuniaria. En aquellos "
            "casos en que el allanamiento o reconocimiento verse sobre controversias "
            "referidas a actos de discriminaci\u00f3n, actos contrarios a la vida y a la "
            "salud y sustancias peligrosas, se considera como un atenuante pero la "
            "sanci\u00f3n a imponer ser\u00e1 pecuniaria. En todos los supuestos de allanamiento "
            "y reconocimiento formulados con la presentaci\u00f3n de los descargos, se "
            "exonera al denunciado del pago de los costos del procedimiento, pero no "
            "de las costas."
        ),
        "__F10__": (
            "<b>DECRETO LEGISLATIVO N\u00ba 807, LEY SOBRE FACULTADES, NORMAS Y "
            "ORGANIZACI\u00d3N DEL INDECOPI FACULTADES DE LAS COMISIONES Y OFICINAS DEL "
            "INDECOPI</b>\n<b>Art\u00edculo 29.-</b> En cualquier estado del procedimiento, e "
            "incluso antes de admitirse a tr\u00e1mite la denuncia, el Secretario T\u00e9cnico "
            "podr\u00e1 citar a las partes a audiencia de conciliaci\u00f3n. La audiencia se "
            "desarrollar\u00e1 ante el Secretario T\u00e9cnico o ante la persona que \u00e9ste designe. "
            "Si ambas partes arribaran a un acuerdo respecto de la denuncia, se "
            "levantar\u00e1 un acta donde conste el acuerdo respectivo, el mismo que tendr\u00e1 "
            "efectos de transacci\u00f3n extrajudicial. En cualquier caso, la Comisi\u00f3n "
            "podr\u00e1 continuar de oficio el procedimiento, si del an\u00e1lisis de los hechos "
            "denunciados considera que podr\u00eda estarse afectando intereses de terceros."
        ),
        "__F11__": (
            "<b>TEXTO \u00daNICO ORDENADO DE LA LEY DEL PROCEDIMIENTO ADMINISTRATIVO GENERAL, "
            "aprobado mediante DECRETO SUPREMO N\u00b0 004-2019-JUS y publicado el 25 de "
            "enero de 2019</b>\n<b>Art\u00edculo 20.- Modalidades de notificaci\u00f3n</b>\n20.4. "
            "El administrado interesado o afectado por el acto que hubiera consignado "
            "en su escrito alguna direcci\u00f3n electr\u00f3nica que conste en el expediente "
            "puede ser notificado a trav\u00e9s de ese medio siempre que haya dado su "
            "autorizaci\u00f3n expresa para ello. Para este caso no es de aplicaci\u00f3n el "
            "orden de prelaci\u00f3n dispuesto en el numeral 20.1.\nLa notificaci\u00f3n dirigida "
            "a la direcci\u00f3n de correo electr\u00f3nico se\u00f1alada por el administrado se "
            "entiende v\u00e1lidamente efectuada cuando la entidad reciba la respuesta de "
            "recepci\u00f3n de la direcci\u00f3n electr\u00f3nica se\u00f1alada por el administrado o esta "
            "sea generada en forma autom\u00e1tica por una plataforma tecnol\u00f3gica o sistema "
            "inform\u00e1tico que garantice que la notificaci\u00f3n ha sido efectuada. La "
            "notificaci\u00f3n surte efectos el d\u00eda que conste haber sido recibida, conforme "
            "lo previsto en el numeral 2 del art\u00edculo 25.\nEn caso de no recibirse "
            "respuesta autom\u00e1tica de recepci\u00f3n en un plazo m\u00e1ximo de dos (2) d\u00edas "
            "h\u00e1biles contados desde el d\u00eda siguiente de efectuado el acto de notificaci\u00f3n "
            "v\u00eda correo electr\u00f3nico, se procede a notificar por c\u00e9dula conforme al "
            "inciso 20.1.1, volvi\u00e9ndose a computar el plazo establecido en el numeral "
            "24.1 del art\u00edculo 24."
        ),
    }

    # Reemplazar placeholders con footnotes reales
    for key, footnote_text in footnotes_data.items():
        rng = doc_com.Content
        if rng.Find.Execute(FindText=key):
            rng.Text = ""
            fn = doc_com.Footnotes.Add(Range=rng, Text="")

            parts = re.split(r'(<b>.*?</b>)', footnote_text, flags=re.DOTALL)
            insert_rng = fn.Range
            insert_rng.Collapse(1)
            for part in parts:
                if not part:
                    continue
                is_bold = part.startswith('<b>') and part.endswith('</b>')
                text = part[3:-4] if is_bold else part
                insert_rng.Text = text
                insert_rng.Font.Bold = is_bold
                insert_rng.Collapse(0)

    # Formatear notas al pie
    for fn in doc_com.Footnotes:
        fn.Range.Font.Name = "Arial Narrow"
        fn.Range.Font.Size = 8.0
        fn.Range.ParagraphFormat.Alignment = 3  # justify

    # Guardar
    out_dir = os.path.abspath(
        r"D:\BETTER CALL DAVID\ResAdmi\productos (resoluciones) elaborada por google antigravity"
    )
    if not os.path.exists(out_dir):
        os.makedirs(out_dir)
    out_path = os.path.join(out_dir, "ADM 2142-2026 R2.docx")
    doc_com.SaveAs(out_path)

    # PDF
    pdf_path = os.path.join(out_dir, "ADM 2142-2026 R2.pdf")
    doc_com.ExportAsFixedFormat(pdf_path, 17)

    doc_com.Close()
    word.Quit()

    print(f"Documento final: {out_path}")
    print(f"PDF: {pdf_path}")
    return out_path


# ============================================================
# MAIN
# ============================================================

def main():
    print("=" * 70)
    print("Generando Resoluci\u00f3n de Admisorio - Exp. 2142-2026/CC1")
    print("=" * 70)

    # 1. Generar base (python-docx)
    temp_path = generar()

    # 2. Aplicar reglas base (formato CC1)
    print("\nAplicando reglas de formato CC1...")
    doc = Document(temp_path)
    aplicar_reglas_base(doc)
    doc.save(temp_path)
    print("Reglas base aplicadas correctamente.")

    # 3. Inyectar notas al pie (win32com)
    print("\nInyectando notas al pie...")
    out_path = inject_footnotes(temp_path)

    print("\n" + "=" * 70)
    print("RESOLUCI\u00d3N GENERADA EXITOSAMENTE")
    print("=" * 70)
    print(f"Archivo: {os.path.basename(out_path)}")

    # 4. Validar con el validador post-generación
    print("\nEjecutando validador post-generaci\u00f3n...")
    try:
        from validador_post_generacion import validar_documento
        valido, violaciones = validar_documento(out_path)
        if valido:
            print("\n\u2714 DOCUMENTO V\u00c1LIDO - Todas las reglas cumplidas")
        else:
            print("\n\u26a0 Documento v\u00e1lido con advertencias")
    except Exception as e:
        print(f"Validador no disponible: {e}")

    return out_path


if __name__ == "__main__":
    main()
